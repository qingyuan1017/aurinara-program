"""Typed, lossless XLSX exchange for the SDTM Specification aggregate."""

import io
import json
from datetime import datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.sdtm.contracts import SdtmSpecification, SdtmTerminologyBinding

_FIXED_TIME = datetime(2000, 1, 1)
_SHEETS = ("Metadata", "Domains", "Variables", "ValueLevel")
_DOMAIN_HEADERS = (
    "code",
    "label",
    "domain_class",
    "structure",
    "keys_json",
    "sort_order_json",
    "expected_output_name",
)
_VARIABLE_HEADERS = (
    "domain",
    "name",
    "label",
    "data_type",
    "length",
    "role",
    "core",
    "origin",
    "derived",
    "terminology_release_id",
    "codelist_code",
    "term_code",
    "extension_revision_id",
)
_VALUE_LEVEL_HEADERS = (
    "domain",
    "id",
    "variable",
    "condition_json",
    "label",
    "data_type",
    "length",
    "terminology_release_id",
    "codelist_code",
    "term_code",
    "extension_revision_id",
)


def _json(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), sort_keys=True)


def _binding_values(
    binding: SdtmTerminologyBinding | None,
) -> tuple[object | None, ...]:
    if binding is None:
        return (None, None, None, None)
    return (
        str(binding.terminology_release_id),
        binding.codelist_code,
        binding.term_code,
        str(binding.extension_revision_id) if binding.extension_revision_id is not None else None,
    )


def _format_sheet(worksheet: Worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, column in enumerate(worksheet.columns, start=1):
        width = min(60, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        worksheet.column_dimensions[get_column_letter(index)].width = width


def render_specification_xlsx(specification: SdtmSpecification) -> bytes:
    """Render every maintained Specification field without inferring clinical content."""

    workbook = Workbook()
    metadata = workbook.active
    if metadata is None:
        raise RuntimeError("new workbook has no active worksheet")
    metadata.title = "Metadata"
    metadata.append(("key", "value"))
    metadata.append(("schema_version", specification.schema_version))
    domains = workbook.create_sheet("Domains")
    variables = workbook.create_sheet("Variables")
    value_level = workbook.create_sheet("ValueLevel")
    domains.append(_DOMAIN_HEADERS)
    variables.append(_VARIABLE_HEADERS)
    value_level.append(_VALUE_LEVEL_HEADERS)
    for domain in specification.domains:
        domains.append(
            (
                domain.code,
                domain.label,
                domain.domain_class,
                domain.structure,
                _json(domain.keys),
                _json(domain.sort_order),
                domain.expected_output_name,
            )
        )
        for variable in domain.variables:
            variables.append(
                (
                    domain.code,
                    variable.name,
                    variable.label,
                    variable.data_type,
                    variable.length,
                    variable.role,
                    variable.core,
                    variable.origin,
                    variable.derived,
                    *_binding_values(variable.terminology_binding),
                )
            )
        for item in domain.value_level_metadata:
            value_level.append(
                (
                    domain.code,
                    item.id,
                    item.variable,
                    _json(item.condition.model_dump(mode="json")),
                    item.label,
                    item.data_type,
                    item.length,
                    *_binding_values(item.terminology_binding),
                )
            )
    for sheet in workbook.worksheets:
        _format_sheet(sheet)
    workbook.properties.creator = "Aurinara"
    workbook.properties.created = _FIXED_TIME
    workbook.properties.modified = _FIXED_TIME
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _records(sheet: Any, headers: tuple[str, ...]) -> list[dict[str, object]]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows or tuple(rows[0]) != headers:
        raise ValueError(f"{sheet.title} sheet headers do not match the versioned contract")
    return [
        {name: value for name, value in zip(headers, row, strict=True)}
        for row in rows[1:]
        if any(value is not None for value in row)
    ]


def _optional(value: object) -> object | None:
    return None if value is None or value == "" else value


def _binding(record: dict[str, object]) -> dict[str, object] | None:
    release = _optional(record["terminology_release_id"])
    codelist = _optional(record["codelist_code"])
    term = _optional(record["term_code"])
    extension = _optional(record["extension_revision_id"])
    if release is None and codelist is None and term is None and extension is None:
        return None
    if release is None or codelist is None:
        raise ValueError("terminology bindings require release and codelist values")
    return {
        "terminology_release_id": release,
        "codelist_code": codelist,
        "term_code": term,
        "extension_revision_id": extension,
    }


def parse_specification_xlsx(data: bytes) -> SdtmSpecification:
    """Validate all workbook sheets before returning one atomic Specification contract."""

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        if tuple(workbook.sheetnames) != _SHEETS:
            raise ValueError(f"workbook sheets must be exactly: {', '.join(_SHEETS)}")
        metadata_rows = _records(workbook["Metadata"], ("key", "value"))
        metadata = {str(row["key"]): row["value"] for row in metadata_rows}
        domain_rows = _records(workbook["Domains"], _DOMAIN_HEADERS)
        variable_rows = _records(workbook["Variables"], _VARIABLE_HEADERS)
        value_rows = _records(workbook["ValueLevel"], _VALUE_LEVEL_HEADERS)
        domain_codes = {str(row["code"]) for row in domain_rows}
        child_domains = {str(row["domain"]) for row in (*variable_rows, *value_rows)}
        if not child_domains <= domain_codes:
            raise ValueError("Variables and ValueLevel rows must reference declared domains")
        domains: list[dict[str, object]] = []
        for domain in domain_rows:
            code = str(domain["code"])
            variables: list[dict[str, object]] = []
            for row in variable_rows:
                if str(row["domain"]) != code:
                    continue
                variables.append(
                    {
                        "name": row["name"],
                        "label": row["label"],
                        "data_type": row["data_type"],
                        "length": _optional(row["length"]),
                        "role": row["role"],
                        "core": row["core"],
                        "origin": row["origin"],
                        "derived": row["derived"],
                        "terminology_binding": _binding(row),
                    }
                )
            value_level: list[dict[str, object]] = []
            for row in value_rows:
                if str(row["domain"]) != code:
                    continue
                value_level.append(
                    {
                        "id": row["id"],
                        "variable": row["variable"],
                        "condition": json.loads(str(row["condition_json"])),
                        "label": row["label"],
                        "data_type": row["data_type"],
                        "length": _optional(row["length"]),
                        "terminology_binding": _binding(row),
                    }
                )
            domains.append(
                {
                    "code": code,
                    "label": domain["label"],
                    "domain_class": domain["domain_class"],
                    "structure": domain["structure"],
                    "keys": json.loads(str(domain["keys_json"])),
                    "sort_order": json.loads(str(domain["sort_order_json"])),
                    "expected_output_name": domain["expected_output_name"],
                    "variables": variables,
                    "value_level_metadata": value_level,
                }
            )
        return SdtmSpecification.model_validate(
            {"schema_version": metadata.get("schema_version"), "domains": domains}
        )
    finally:
        workbook.close()
