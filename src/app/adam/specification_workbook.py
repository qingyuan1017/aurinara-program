"""Typed, lossless XLSX exchange for the ADaM Specification aggregate."""

import io
import json
from datetime import datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.adam.contracts import AdamSpecification, AdamTerminologyBinding
from app.common.zip_normalization import normalize_zip_bytes

_FIXED_TIME = datetime(2000, 1, 1)
_SHEETS = (
    "Metadata",
    "Datasets",
    "Variables",
    "Parameters",
    "ValueLevel",
    "Methods",
    "TraceTargets",
)
_DATASET_HEADERS = (
    "id",
    "code",
    "label",
    "dataset_class",
    "structure",
    "keys_json",
    "sort_order_json",
    "purpose",
    "expected_output_name",
)
_VARIABLE_HEADERS = (
    "dataset",
    "id",
    "name",
    "label",
    "data_type",
    "length",
    "role",
    "origin_json",
    "method_id",
    "output_node_id",
    "terminology_release_id",
    "codelist_code",
    "term_code",
    "extension_revision_id",
)
_PARAMETER_HEADERS = (
    "dataset",
    "id",
    "code",
    "label",
    "variable",
    "value_json",
    "terminology_release_id",
    "codelist_code",
    "term_code",
    "extension_revision_id",
)
_VALUE_LEVEL_HEADERS = (
    "dataset",
    "id",
    "variable",
    "condition_json",
    "label",
    "data_type",
    "length",
    "origin_json",
    "method_id",
    "terminology_release_id",
    "codelist_code",
    "term_code",
    "extension_revision_id",
)
_METHOD_HEADERS = ("id", "name", "description", "plan_node_ids_json")
_TRACE_TARGET_HEADERS = (
    "relationship_type",
    "source_module",
    "source_entity_type",
    "source_entity_id",
    "target_module",
    "target_entity_type",
    "target_entity_id",
)


def _json(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), sort_keys=True)


def _binding_values(
    binding: AdamTerminologyBinding | None,
) -> tuple[object | None, ...]:
    if binding is None:
        return (None, None, None, None)
    return (
        str(binding.terminology_release_id),
        binding.codelist_code,
        binding.term_code,
        str(binding.extension_revision_id) if binding.extension_revision_id is not None else None,
    )


def _format_sheet(worksheet: Worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, column in enumerate(worksheet.columns, start=1):
        width = min(60, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        worksheet.column_dimensions[get_column_letter(index)].width = width


def render_specification_xlsx(specification: AdamSpecification) -> bytes:
    """Render every maintained ADaM Specification field without clinical inference."""

    workbook = Workbook()
    metadata = workbook.active
    if metadata is None:
        raise RuntimeError("new workbook has no active worksheet")
    metadata.title = "Metadata"
    metadata.append(("key", "value"))
    metadata.append(("schema_version", specification.schema_version))
    datasets = workbook.create_sheet("Datasets")
    variables = workbook.create_sheet("Variables")
    parameters = workbook.create_sheet("Parameters")
    value_level = workbook.create_sheet("ValueLevel")
    methods = workbook.create_sheet("Methods")
    trace_targets = workbook.create_sheet("TraceTargets")
    datasets.append(_DATASET_HEADERS)
    variables.append(_VARIABLE_HEADERS)
    parameters.append(_PARAMETER_HEADERS)
    value_level.append(_VALUE_LEVEL_HEADERS)
    methods.append(_METHOD_HEADERS)
    trace_targets.append(_TRACE_TARGET_HEADERS)
    for dataset in specification.datasets:
        datasets.append(
            (
                str(dataset.id),
                dataset.code,
                dataset.label,
                dataset.dataset_class,
                dataset.structure,
                _json(dataset.keys),
                _json(dataset.sort_order),
                dataset.purpose,
                dataset.expected_output_name,
            )
        )
        for variable in dataset.variables:
            variables.append(
                (
                    dataset.code,
                    str(variable.id),
                    variable.name,
                    variable.label,
                    variable.data_type,
                    variable.length,
                    variable.role,
                    _json(variable.origin.model_dump(mode="json")),
                    variable.method_id,
                    variable.output_node_id,
                    *_binding_values(variable.terminology_binding),
                )
            )
        for parameter in dataset.parameters:
            parameters.append(
                (
                    dataset.code,
                    str(parameter.id),
                    parameter.code,
                    parameter.label,
                    parameter.variable,
                    _json(parameter.value),
                    *_binding_values(parameter.terminology_binding),
                )
            )
        for item in dataset.value_level_metadata:
            value_level.append(
                (
                    dataset.code,
                    str(item.id),
                    item.variable,
                    _json(item.condition.model_dump(mode="json")),
                    item.label,
                    item.data_type,
                    item.length,
                    _json(item.origin.model_dump(mode="json")),
                    item.method_id,
                    *_binding_values(item.terminology_binding),
                )
            )
    for method in specification.methods:
        methods.append((method.id, method.name, method.description, _json(method.plan_node_ids)))
    for trace_target in specification.trace_targets:
        trace_targets.append(
            (
                trace_target.relationship_type,
                trace_target.source.module,
                trace_target.source.entity_type,
                str(trace_target.source.entity_id),
                trace_target.target.module,
                trace_target.target.entity_type,
                str(trace_target.target.entity_id),
            )
        )
    for sheet in workbook.worksheets:
        _format_sheet(sheet)
    workbook.properties.creator = "Aurinara"
    workbook.properties.created = _FIXED_TIME
    workbook.properties.modified = _FIXED_TIME
    output = io.BytesIO()
    workbook.save(output)
    return normalize_zip_bytes(output.getvalue())


def _records(sheet: Any, headers: tuple[str, ...]) -> list[dict[str, object]]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows or tuple(rows[0]) != headers:
        raise ValueError(f"{sheet.title} sheet headers do not match the versioned contract")
    return [
        {name: value for name, value in zip(headers, row, strict=True)}
        for row in rows[1:]
        if any(value is not None for value in row)
    ]


def _optional(value: object) -> object | None:
    return None if value is None or value == "" else value


def _binding(record: dict[str, object]) -> dict[str, object] | None:
    release = _optional(record["terminology_release_id"])
    codelist = _optional(record["codelist_code"])
    term = _optional(record["term_code"])
    extension = _optional(record["extension_revision_id"])
    if release is None and codelist is None and term is None and extension is None:
        return None
    if release is None or codelist is None:
        raise ValueError("terminology bindings require release and codelist values")
    return {
        "terminology_release_id": release,
        "codelist_code": codelist,
        "term_code": term,
        "extension_revision_id": extension,
    }


def parse_specification_xlsx(data: bytes) -> AdamSpecification:
    """Validate every workbook sheet before returning one atomic Specification."""

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        if tuple(workbook.sheetnames) != _SHEETS:
            raise ValueError(f"workbook sheets must be exactly: {', '.join(_SHEETS)}")
        metadata_rows = _records(workbook["Metadata"], ("key", "value"))
        metadata = {str(row["key"]): row["value"] for row in metadata_rows}
        dataset_rows = _records(workbook["Datasets"], _DATASET_HEADERS)
        variable_rows = _records(workbook["Variables"], _VARIABLE_HEADERS)
        parameter_rows = _records(workbook["Parameters"], _PARAMETER_HEADERS)
        value_rows = _records(workbook["ValueLevel"], _VALUE_LEVEL_HEADERS)
        method_rows = _records(workbook["Methods"], _METHOD_HEADERS)
        trace_rows = _records(workbook["TraceTargets"], _TRACE_TARGET_HEADERS)
        dataset_codes = {str(row["code"]) for row in dataset_rows}
        child_datasets = {
            str(row["dataset"]) for row in (*variable_rows, *parameter_rows, *value_rows)
        }
        if not child_datasets <= dataset_codes:
            raise ValueError("child rows must reference declared ADaM datasets")
        datasets: list[dict[str, object]] = []
        for dataset in dataset_rows:
            code = str(dataset["code"])
            variables = [
                {
                    "id": row["id"],
                    "name": row["name"],
                    "label": row["label"],
                    "data_type": row["data_type"],
                    "length": _optional(row["length"]),
                    "role": row["role"],
                    "origin": json.loads(str(row["origin_json"])),
                    "method_id": _optional(row["method_id"]),
                    "output_node_id": row["output_node_id"],
                    "terminology_binding": _binding(row),
                }
                for row in variable_rows
                if str(row["dataset"]) == code
            ]
            parameters = [
                {
                    "id": row["id"],
                    "code": row["code"],
                    "label": row["label"],
                    "variable": row["variable"],
                    "value": json.loads(str(row["value_json"])),
                    "terminology_binding": _binding(row),
                }
                for row in parameter_rows
                if str(row["dataset"]) == code
            ]
            value_level = [
                {
                    "id": row["id"],
                    "variable": row["variable"],
                    "condition": json.loads(str(row["condition_json"])),
                    "label": row["label"],
                    "data_type": row["data_type"],
                    "length": _optional(row["length"]),
                    "origin": json.loads(str(row["origin_json"])),
                    "method_id": _optional(row["method_id"]),
                    "terminology_binding": _binding(row),
                }
                for row in value_rows
                if str(row["dataset"]) == code
            ]
            datasets.append(
                {
                    "id": dataset["id"],
                    "code": code,
                    "label": dataset["label"],
                    "dataset_class": dataset["dataset_class"],
                    "structure": dataset["structure"],
                    "keys": json.loads(str(dataset["keys_json"])),
                    "sort_order": json.loads(str(dataset["sort_order_json"])),
                    "purpose": dataset["purpose"],
                    "expected_output_name": dataset["expected_output_name"],
                    "variables": variables,
                    "parameters": parameters,
                    "value_level_metadata": value_level,
                }
            )
        methods = [
            {
                "id": row["id"],
                "name": row["name"],
                "description": row["description"],
                "plan_node_ids": json.loads(str(row["plan_node_ids_json"])),
            }
            for row in method_rows
        ]
        trace_targets = [
            {
                "relationship_type": row["relationship_type"],
                "source": {
                    "module": row["source_module"],
                    "entity_type": row["source_entity_type"],
                    "entity_id": row["source_entity_id"],
                },
                "target": {
                    "module": row["target_module"],
                    "entity_type": row["target_entity_type"],
                    "entity_id": row["target_entity_id"],
                },
            }
            for row in trace_rows
        ]
        return AdamSpecification.model_validate(
            {
                "schema_version": metadata.get("schema_version"),
                "datasets": datasets,
                "methods": methods,
                "trace_targets": trace_targets,
            }
        )
    finally:
        workbook.close()
