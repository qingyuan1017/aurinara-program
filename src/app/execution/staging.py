"""Technical tabular input normalization for isolated execution workspaces."""

import csv
import io
import json
from pathlib import PurePath
from typing import cast

import pandas as pd
import pyreadstat
from openpyxl import load_workbook


def tabular_csv_bytes(data: bytes, filename: str) -> bytes:
    """Convert one supported immutable tabular artifact into deterministic UTF-8 CSV."""

    suffix = PurePath(filename).suffix.lower()
    if suffix == ".csv":
        _require_csv_headers(data)
        return data
    if suffix == ".xpt":
        headers, rows = _xpt_rows(data)
    elif suffix == ".json":
        headers, rows = _json_rows(data)
    elif suffix == ".xlsx":
        headers, rows = _xlsx_rows(data)
    else:
        raise ValueError(f"execution input format '{suffix}' is unsupported")
    _require_headers(headers)
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(headers)
    writer.writerows(rows)
    return output.getvalue().encode("utf-8")


def _require_csv_headers(data: bytes) -> None:
    try:
        reader = csv.reader(io.StringIO(data.decode("utf-8-sig")))
        headers = next(reader, [])
    except (UnicodeError, csv.Error) as error:
        raise ValueError("execution CSV input must be valid UTF-8 CSV") from error
    _require_headers(headers)


def _require_headers(headers: list[str]) -> None:
    if not headers or any(not item for item in headers) or len(headers) != len(set(headers)):
        raise ValueError("execution input columns must be present and unique")


def _json_rows(data: bytes) -> tuple[list[str], list[list[object]]]:
    try:
        payload = json.loads(data)
    except (UnicodeError, json.JSONDecodeError) as error:
        raise ValueError("execution JSON input must be valid UTF-8 JSON") from error
    if not isinstance(payload, list) or any(not isinstance(item, dict) for item in payload):
        raise ValueError("execution JSON input must be an array of flat objects")
    headers: list[str] = []
    for item in payload:
        for key, value in item.items():
            if not isinstance(key, str) or isinstance(value, (dict, list)):
                raise ValueError("execution JSON keys must be strings and values must be scalar")
            if key not in headers:
                headers.append(key)
    return headers, [[item.get(header) for header in headers] for item in payload]


def _xlsx_rows(data: bytes) -> tuple[list[str], list[list[object]]]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (OSError, ValueError) as error:
        raise ValueError("execution XLSX input is invalid") from error
    try:
        sheet = workbook.active
        if sheet is None:
            raise ValueError("execution XLSX input has no active worksheet")
        values = sheet.iter_rows(values_only=True)
        headers = [str(value) if value is not None else "" for value in next(values, ())]
        rows = [[cast(object, value) for value in row] for row in values]
        return headers, rows
    finally:
        workbook.close()


def _xpt_rows(data: bytes) -> tuple[list[str], list[list[object]]]:
    try:
        frame, metadata = pyreadstat.read_xport(io.BytesIO(data), output_format="pandas")
    except pyreadstat.ReadstatError as error:
        raise ValueError("execution XPT input is invalid") from error
    headers = [str(value) for value in metadata.column_names]
    normalized = frame.astype(object).where(pd.notna(frame), None)
    return headers, [list(row) for row in normalized.itertuples(index=False, name=None)]
