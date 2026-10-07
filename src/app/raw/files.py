"""Bounded readers for factual metadata and preview of supported Raw source files."""

import csv
import io
import json
from datetime import date, datetime
from pathlib import PurePath
from typing import BinaryIO

from openpyxl import load_workbook

from app.common.problems import ProblemException
from app.raw.contracts import RawPhysicalColumn, RawPhysicalSnapshot

SUPPORTED_SUFFIXES = frozenset({".csv", ".json", ".xlsx"})


def _unsupported(filename: str) -> ProblemException:
    return ProblemException(
        status_code=422,
        code="raw_file_format_unsupported",
        title="Raw file format is unsupported",
        detail=f"'{PurePath(filename).suffix.lower()}' is not a supported Raw input format.",
    )


def _type_of(value: object) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "number"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    text = str(value)
    try:
        int(text)
        return "integer"
    except ValueError:
        pass
    try:
        float(text)
        return "number"
    except ValueError:
        return "string"


def _merge_type(current: str | None, observed: str | None) -> str | None:
    if observed is None or current == observed:
        return current or observed
    if current is None:
        return observed
    if {current, observed} <= {"integer", "number"}:
        return "number"
    return "string"


def _snapshot(headers: list[str], rows: list[list[object]]) -> RawPhysicalSnapshot:
    types: list[str | None] = [None] * len(headers)
    lengths: list[int] = [0] * len(headers)
    precision: list[int] = [0] * len(headers)
    for row in rows:
        for index, value in enumerate(row[: len(headers)]):
            types[index] = _merge_type(types[index], _type_of(value))
            if value is not None:
                text = str(value)
                lengths[index] = max(lengths[index], len(text))
                if "." in text and text.replace(".", "", 1).replace("-", "", 1).isdigit():
                    precision[index] = max(precision[index], len(text.rsplit(".", maxsplit=1)[1]))
    return RawPhysicalSnapshot(
        columns=tuple(
            RawPhysicalColumn(
                name=name,
                physical_type=types[index] or "string",
                length=lengths[index] or None,
                precision=precision[index] if types[index] == "number" else None,
                ordinal=index + 1,
            )
            for index, name in enumerate(headers)
        ),
        row_count=len(rows),
    )


def _csv_rows(stream: BinaryIO) -> tuple[list[str], list[list[object]]]:
    text = io.TextIOWrapper(stream, encoding="utf-8-sig", newline="")
    reader = csv.reader(text)
    try:
        headers = next(reader)
    except StopIteration:
        headers = []
    return headers, [list(row) for row in reader]


def _json_rows(stream: BinaryIO) -> tuple[list[str], list[list[object]]]:
    payload = json.load(stream)
    if not isinstance(payload, list) or any(not isinstance(item, dict) for item in payload):
        raise ProblemException(
            status_code=422,
            code="raw_json_shape_invalid",
            title="Raw JSON shape is invalid",
            detail="Raw JSON must be an array of flat objects.",
        )
    headers: list[str] = []
    for item in payload:
        for key, value in item.items():
            if not isinstance(key, str) or isinstance(value, (dict, list)):
                raise ProblemException(
                    status_code=422,
                    code="raw_json_shape_invalid",
                    title="Raw JSON shape is invalid",
                    detail="Raw JSON keys must be strings and values must be scalar.",
                )
            if key not in headers:
                headers.append(key)
    return headers, [[item.get(header) for header in headers] for item in payload]


def _xlsx_rows(stream: BinaryIO) -> tuple[list[str], list[list[object]]]:
    workbook = load_workbook(stream, read_only=True, data_only=True)
    sheet = workbook.active
    if sheet is None:
        raise ProblemException(
            status_code=422,
            code="raw_physical_schema_invalid",
            title="Raw physical schema is invalid",
            detail="The workbook has no active worksheet.",
        )
    values = sheet.iter_rows(values_only=True)
    first = next(values, ())
    headers = [str(value) if value is not None else "" for value in first]
    rows: list[list[object]] = []
    for row in values:
        rows.append(list(row))
    workbook.close()
    return headers, rows


def read_rows(stream: BinaryIO, filename: str) -> tuple[list[str], list[list[object]]]:
    """Read supported bytes without retaining them after this request boundary."""

    suffix = PurePath(filename).suffix.lower()
    if suffix == ".csv":
        return _csv_rows(stream)
    if suffix == ".json":
        return _json_rows(stream)
    if suffix == ".xlsx":
        return _xlsx_rows(stream)
    raise _unsupported(filename)


def inspect_file(stream: BinaryIO, filename: str) -> RawPhysicalSnapshot:
    headers, rows = read_rows(stream, filename)
    if not headers or any(not header for header in headers) or len(headers) != len(set(headers)):
        raise ProblemException(
            status_code=422,
            code="raw_physical_schema_invalid",
            title="Raw physical schema is invalid",
            detail="Column names must be present and unique.",
        )
    return _snapshot(headers, rows)


def preview_rows(
    stream: BinaryIO,
    filename: str,
    columns: tuple[str, ...],
    offset: int,
    limit: int,
) -> tuple[tuple[str, ...], tuple[dict[str, object | None], ...], int]:
    """Return only the requested bounded row window."""

    headers, rows = read_rows(stream, filename)
    selected = list(columns) if columns else headers
    unknown = sorted(set(selected) - set(headers))
    if unknown:
        raise ProblemException(
            status_code=422,
            code="raw_preview_column_unknown",
            title="Raw preview column is unknown",
            detail=f"Unknown columns: {', '.join(unknown)}",
        )
    positions = {name: index for index, name in enumerate(headers)}
    window = rows[offset : offset + limit]
    result = tuple(
        {name: row[positions[name]] if positions[name] < len(row) else None for name in selected}
        for row in window
    )
    return tuple(selected), result, len(rows)
