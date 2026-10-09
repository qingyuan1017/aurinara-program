"""PostgreSQL and storage-backed Raw-to-SDTM authoring workflow."""

import io
import os
import subprocess
import sys
from collections.abc import AsyncIterator
from uuid import UUID, uuid4

import pyreadstat
import pytest
from docx import Document as DocxDocument
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.pdfgen.canvas import Canvas
from sqlalchemy import text

from app.artifacts.builds import resolve_build_input
from app.catalog_seed import seed
from app.common.canonical import fingerprint
from app.database import create_database_engine, create_session_factory
from app.execution.adapters import ExecutionRequest, ExecutionResult
from app.execution.jobs import JobService
from app.execution.worker import process_execution_job
from app.main import create_app
from app.sdtm.define_xml import validate_define_xml
from tests.conftest import make_test_settings
from tests.integration.test_phase1_core import MemoryStorage, TestTokenVerifier
from workers.job_worker import _process_job

pytestmark = pytest.mark.skipif(
    os.getenv("AURINARA_INTEGRATION") != "1",
    reason="set AURINARA_INTEGRATION=1 with the Compose PostgreSQL service running",
)


class HostPythonExecutor:
    """Test-only adapter exercising the generated program without weakening production policy."""

    def execute(self, request: ExecutionRequest) -> ExecutionResult:
        completed = subprocess.run(
            [sys.executable, "-I", str(request.workspace / "program.py")],
            check=False,
            capture_output=True,
            env={**os.environ, "AURINARA_WORKSPACE": str(request.workspace)},
            timeout=request.limits.runtime_seconds,
        )
        return ExecutionResult(
            status="succeeded" if completed.returncode == 0 else "failed",
            exit_code=completed.returncode,
            stdout=completed.stdout,
            stderr=completed.stderr,
            failure=None if completed.returncode == 0 else {"code": "program_failed"},
        )


@pytest.fixture
async def phase3_client() -> AsyncIterator[AsyncClient]:
    settings = make_test_settings()
    engine = create_database_engine(settings)
    factory = create_session_factory(engine)
    with engine.begin() as connection:
        connection.execute(
            text(
                "TRUNCATE execution_runs, sdtm_csdrg_revisions, sdtm_define_revisions, "
                "sdtm_acrf_annotations, "
                "sdtm_derivation_plan_revisions, "
                "sdtm_specification_revisions, "
                "sdtm_input_selections, raw_variables, raw_dataset_file_bindings, raw_datasets, "
                "protocol_evidence, protocol_usdm_documents, document_contents, "
                "document_dependency_selections, documents, integration_delivery_attempts, "
                "integration_subscriptions, integration_secret_versions, external_identifiers, "
                "integrations, reconciliation_runs, validation_runs, source_files, artifacts, "
                "jobs, idempotency_records, outbox_events, study_memberships, studies CASCADE"
            )
        )
    storage = MemoryStorage()
    with factory() as session:
        seed(session, storage)
    app = create_app(
        settings=settings,
        session_factory=factory,
        object_storage=storage,
        token_verifier=TestTokenVerifier(),
    )
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://testserver"
    ) as client:
        yield client
    engine.dispose()


@pytest.mark.anyio
async def test_sdtm_specification_xlsx_export_and_idempotent_import(
    phase3_client: AsyncClient,
) -> None:
    client = phase3_client
    headers = {"Authorization": "Bearer test-token"}
    study = await client.post(
        "/api/v1/studies", headers=headers, json={"code": "P3-XLSX", "name": "Phase 3 XLSX"}
    )
    sdtm = await client.post(
        f"/api/v1/studies/{study.json()['id']}/sdtms",
        headers=headers,
        json={"code": "SDTM", "name": "SDTM"},
    )
    payload = {
        "domains": [
            {
                "code": "DM",
                "label": "Demographics",
                "domain_class": "Special Purpose",
                "structure": "One record per subject",
                "keys": ["USUBJID"],
                "sort_order": ["USUBJID"],
                "expected_output_name": "dm.xpt",
                "variables": [
                    {
                        "name": "USUBJID",
                        "label": "Unique Subject Identifier",
                        "data_type": "character",
                        "length": 40,
                        "role": "Identifier",
                        "core": "Req",
                        "origin": "Derived",
                        "derived": True,
                    }
                ],
            }
        ]
    }
    specification = await client.put(
        f"/api/v1/sdtms/{sdtm.json()['id']}/specification",
        headers={**headers, "If-Match": sdtm.headers["etag"]},
        json=payload,
    )
    assert specification.status_code == 200, specification.text
    exported = await client.get(
        f"/api/v1/sdtms/{sdtm.json()['id']}/specification/export?format=xlsx",
        headers=headers,
    )
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(io.BytesIO(exported.content))
    workbook["Domains"]["B2"] = "Manually maintained label"
    changed = io.BytesIO()
    workbook.save(changed)
    workbook.close()
    import_headers = {
        **headers,
        "If-Match": specification.headers["etag"],
        "Idempotency-Key": "sdtm-specification-xlsx-1",
        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }
    imported = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/specification/import",
        headers=import_headers,
        content=changed.getvalue(),
    )
    assert imported.status_code == 200, imported.text
    assert imported.json()["revision"] == 2
    assert imported.json()["payload"]["domains"][0]["label"] == "Manually maintained label"
    replay = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/specification/import",
        headers=import_headers,
        content=changed.getvalue(),
    )
    assert replay.status_code == 200, replay.text
    assert replay.json()["id"] == imported.json()["id"]


@pytest.mark.anyio
async def test_raw_latest_validation_run_resolution(
    phase3_client: AsyncClient,
) -> None:
    client = phase3_client
    headers = {"Authorization": "Bearer test-token"}
    study = await client.post(
        "/api/v1/studies", headers=headers, json={"code": "P3-LVR", "name": "Phase 3 LVR"}
    )
    study_id = study.json()["id"]
    source = await client.post(
        f"/api/v1/studies/{study_id}/source-files",
        headers=headers,
        files={
            "file": (
                "demog.json",
                b'[{"SUBJID":"01","AGE":42},{"SUBJID":"02","AGE":38}]',
                "application/json",
            )
        },
    )
    raw = await client.post(
        f"/api/v1/studies/{study_id}/raw",
        headers=headers,
        json={"code": "RAW", "name": "Raw"},
    )
    raw_id = raw.json()["id"]

    # Before any dataset or validation run exists, resolution reports no run.
    empty = await client.get(f"/api/v1/raw/{raw_id}/latest-validation-run", headers=headers)
    assert empty.status_code == 404, empty.text
    assert empty.json()["code"] == "raw_validation_run_not_found"

    dataset = await client.post(
        f"/api/v1/raw/{raw_id}/datasets",
        headers=headers,
        json={"code": "DEMOG", "label": "Demographics", "source_file_id": source.json()["id"]},
    )
    assert dataset.status_code == 201, dataset.text
    initialized = await client.post(
        f"/api/v1/raw-datasets/{dataset.json()['id']}/variables/initialize", headers=headers
    )
    assert initialized.status_code == 200, initialized.text

    # A failed run must not be resolved as the latest successful run: remove a
    # maintained variable so validation fails first.
    variable_id = initialized.json()["variables"][0]["id"]
    deleted = await client.delete(
        f"/api/v1/raw-variables/{variable_id}",
        headers={**headers, "If-Match": '"revision-1"'},
    )
    assert deleted.status_code == 204, deleted.text
    failed_run = await client.post(f"/api/v1/raw/{raw_id}/validate", headers=headers)
    assert failed_run.json()["status"] == "failed", failed_run.text
    still_empty = await client.get(
        f"/api/v1/raw/{raw_id}/latest-validation-run", headers=headers
    )
    assert still_empty.status_code == 404, still_empty.text

    # Re-fill the physical metadata and pass validation.
    reinitialized = await client.post(
        f"/api/v1/raw-datasets/{dataset.json()['id']}/variables/initialize", headers=headers
    )
    assert reinitialized.status_code == 200, reinitialized.text
    first = await client.post(f"/api/v1/raw/{raw_id}/validate", headers=headers)
    assert first.json()["status"] == "succeeded", first.text
    current_raw = await client.get(f"/api/v1/raw/{raw_id}", headers=headers)
    second = await client.post(f"/api/v1/raw/{raw_id}/validate", headers=headers)
    assert second.json()["status"] == "succeeded", second.text

    resolved = await client.get(
        f"/api/v1/raw/{raw_id}/latest-validation-run", headers=headers
    )
    assert resolved.status_code == 200, resolved.text
    body = resolved.json()
    # The newest succeeded run is resolved (not the first), and the advisory
    # raw_revision matches the run's pinned Raw revision.
    assert body["validation_run_id"] == second.json()["id"]
    assert body["validation_run_id"] != first.json()["id"]
    assert body["raw_revision"] == current_raw.json()["revision"]
    assert len(body["input_fingerprint"]) == 64

    # Study access is enforced: an unknown raw id is a 404.
    unknown = await client.get(
        f"/api/v1/raw/{uuid4()}/latest-validation-run", headers=headers
    )
    assert unknown.status_code == 404, unknown.text
    assert unknown.json()["code"] == "raw_not_found"


@pytest.mark.anyio
async def test_raw_to_aligned_sdtm_and_upstream_staleness(
    phase3_client: AsyncClient,
) -> None:
    client = phase3_client
    headers = {"Authorization": "Bearer test-token"}
    study = await client.post(
        "/api/v1/studies", headers=headers, json={"code": "P3-001", "name": "Phase 3"}
    )
    study_id = study.json()["id"]
    source = await client.post(
        f"/api/v1/studies/{study_id}/source-files",
        headers=headers,
        files={
            "file": (
                "demog.json",
                b'[{"SUBJID":"01","AGE":42},{"SUBJID":"02","AGE":38}]',
                "application/json",
            )
        },
    )
    raw = await client.post(
        f"/api/v1/studies/{study_id}/raw",
        headers=headers,
        json={"code": "RAW", "name": "Raw"},
    )
    dataset = await client.post(
        f"/api/v1/raw/{raw.json()['id']}/datasets",
        headers=headers,
        json={"code": "DEMOG", "label": "Demographics", "source_file_id": source.json()["id"]},
    )
    assert dataset.status_code == 201, dataset.text
    initialized = await client.post(
        f"/api/v1/raw-datasets/{dataset.json()['id']}/variables/initialize", headers=headers
    )
    assert initialized.status_code == 200, initialized.text
    preview = await client.post(
        f"/api/v1/raw-datasets/{dataset.json()['id']}/preview",
        headers=headers,
        json={"columns": ["SUBJID"], "limit": 1},
    )
    assert preview.json()["rows"] == [{"SUBJID": "01"}]
    raw_validation = await client.post(f"/api/v1/raw/{raw.json()['id']}/validate", headers=headers)
    assert raw_validation.json()["status"] == "succeeded", raw_validation.text

    current_raw = await client.get(f"/api/v1/raw/{raw.json()['id']}", headers=headers)
    raw_reconciliation = await client.post(
        f"/api/v1/raw/{raw.json()['id']}/reconcile",
        headers={**headers, "Idempotency-Key": "reconcile-raw-schema-1"},
        json={"expected_raw_revision": current_raw.json()["revision"]},
    )
    assert raw_reconciliation.status_code == 202, raw_reconciliation.text
    application = client._transport.app  # type: ignore[attr-defined]
    factory = application.state.session_factory
    with factory.begin() as session:
        raw_reconciliation_job = JobService().claim(
            session, worker_id="phase3-worker", lease_seconds=60
        )
        assert raw_reconciliation_job is not None
    _process_job(
        factory,
        application.state.module_registry,
        application.state.object_storage,
        raw_reconciliation_job,
    )
    reconciled_raw = await client.get(
        f"/api/v1/reconciliation-runs/{raw_reconciliation.json()['id']}", headers=headers
    )
    assert reconciled_raw.json()["status"] == "succeeded", reconciled_raw.text
    assert reconciled_raw.json()["error_count"] == 0
    raw_report = await client.get(
        f"/api/v1/artifacts/{reconciled_raw.json()['report_artifact_ids'][0]}/download",
        headers=headers,
    )
    assert "_object_key" not in raw_report.text
    assert '"SUBJID"' in raw_report.text
    assert '"01"' not in raw_report.text

    # The Artifacts panel resolves the exact build-input fingerprint from the
    # preview endpoint and sends it back as expected_input_fingerprint, so the
    # build must be ACCEPTED (not 409 stale_input_fingerprint).
    raw_artifact_types = ("raw_metadata_json", "raw_data_dictionary_xlsx")
    raw_fingerprint = await client.post(
        f"/api/v1/raw/{raw.json()['id']}/artifacts/fingerprint",
        headers=headers,
        json={"artifact_types": list(raw_artifact_types)},
    )
    assert raw_fingerprint.status_code == 200, raw_fingerprint.text
    raw_build = await client.post(
        f"/api/v1/raw/{raw.json()['id']}/artifacts",
        headers={**headers, "Idempotency-Key": "build-raw-metadata-1"},
        json={
            "artifact_types": list(raw_artifact_types),
            "expected_input_fingerprint": raw_fingerprint.json()["input_fingerprint"],
        },
    )
    assert raw_build.status_code == 202, raw_build.text
    with factory.begin() as session:
        raw_build_job = JobService().claim(session, worker_id="phase3-worker", lease_seconds=60)
        assert raw_build_job is not None
    _process_job(
        factory,
        application.state.module_registry,
        application.state.object_storage,
        raw_build_job,
    )
    completed_raw_build = await client.get(
        f"/api/v1/artifact-builds/{raw_build.json()['id']}", headers=headers
    )
    assert completed_raw_build.json()["status"] == "succeeded", completed_raw_build.text
    assert {
        (await client.get(f"/api/v1/artifacts/{artifact_id}", headers=headers)).json()[
            "artifact_type"
        ]
        for artifact_id in completed_raw_build.json()["result_artifact_ids"]
    } == set(raw_artifact_types)

    raw_dependencies = await client.get(
        f"/api/v1/raw/{raw.json()['id']}/dependencies", headers=headers
    )
    sdtm = await client.post(
        f"/api/v1/studies/{study_id}/sdtms",
        headers=headers,
        json={"code": "SDTM", "name": "SDTM"},
    )
    sdtm_dependencies = await client.put(
        f"/api/v1/sdtms/{sdtm.json()['id']}/dependencies",
        headers={**headers, "If-Match": sdtm.headers["etag"]},
        json={
            "templates": [
                {
                    "purpose": "presentation",
                    "artifact_type": artifact_type,
                    "template_version_id": "30000000-0000-0000-0000-000000000004",
                }
                for artifact_type in ("sdtm_csdrg_docx", "sdtm_csdrg_pdf")
            ]
        },
    )
    assert sdtm_dependencies.status_code == 200, sdtm_dependencies.text
    selection = await client.put(
        f"/api/v1/sdtms/{sdtm.json()['id']}/input-selection",
        headers={**headers, "If-Match": sdtm_dependencies.headers["etag"]},
        json={
            "raw_id": raw.json()["id"],
            "raw_revision": current_raw.json()["revision"],
            "validation_run_id": raw_validation.json()["id"],
            "source_artifacts": [
                {
                    "dataset_code": "DEMOG",
                    "artifact_id": source.json()["artifact_id"],
                    "checksum": source.json()["sha256"],
                }
            ],
            "raw_dependency_fingerprint": raw_dependencies.json()["fingerprint"],
        },
    )
    assert selection.status_code == 200, selection.text
    specification = await client.put(
        f"/api/v1/sdtms/{sdtm.json()['id']}/specification",
        headers={**headers, "If-Match": selection.headers["etag"]},
        json={
            "domains": [
                {
                    "code": "DM",
                    "label": "Demographics",
                    "domain_class": "Special Purpose",
                    "structure": "One record per subject",
                    "keys": ["STUDYID", "USUBJID"],
                    "sort_order": ["STUDYID", "USUBJID"],
                    "expected_output_name": "dm.xpt",
                    "variables": [
                        {
                            "name": "STUDYID",
                            "label": "Study Identifier",
                            "data_type": "character",
                            "length": 20,
                            "role": "Identifier",
                            "core": "Req",
                            "origin": "Assigned",
                            "derived": True,
                        },
                        {
                            "name": "USUBJID",
                            "label": "Unique Subject Identifier",
                            "data_type": "character",
                            "length": 40,
                            "role": "Identifier",
                            "core": "Req",
                            "origin": "Derived",
                            "derived": True,
                        },
                    ],
                }
            ]
        },
    )
    assert specification.status_code == 200, specification.text
    plan = await client.put(
        f"/api/v1/sdtms/{sdtm.json()['id']}/derivation-plan",
        headers={**headers, "If-Match": specification.headers["etag"]},
        json={
            "specification_revision": 1,
            "input_selection_revision": 1,
            "nodes": [
                {"node_type": "raw_read", "id": "read", "raw_dataset_code": "DEMOG"},
                {
                    "node_type": "derive",
                    "id": "derive",
                    "target_domain": "DM",
                    "assignments": {
                        "STUDYID": {"kind": "literal", "value": "P3-001"},
                        "USUBJID": {"kind": "column", "dataset": "DEMOG", "column": "SUBJID"},
                    },
                },
                {"node_type": "domain_write", "id": "write", "domain": "DM"},
            ],
            "edges": [
                {"source_node": "read", "target_node": "derive"},
                {"source_node": "derive", "target_node": "write"},
            ],
        },
    )
    assert plan.status_code == 200, plan.text
    plan_preview = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/derivation-plan/preview",
        headers=headers,
        json={"domains": ["DM"], "limit": 1},
    )
    assert plan_preview.status_code == 200, plan_preview.text
    assert plan_preview.json()["outputs"]["DM"] == [
        {"SUBJID": "01", "AGE": 42, "STUDYID": "P3-001", "USUBJID": "01"}
    ]

    pdf_stream = io.BytesIO()
    pdf = Canvas(pdf_stream)
    pdf.drawString(72, 720, "Synthetic demographics CRF")
    pdf.save()
    crf = await client.post(
        f"/api/v1/studies/{study_id}/source-files",
        headers=headers,
        files={"file": ("crf.pdf", pdf_stream.getvalue(), "application/pdf")},
    )
    raw_variable_id = initialized.json()["variables"][0]["id"]
    annotation = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/acrf/annotations",
        headers={**headers, "If-Match": plan.headers["etag"]},
        json={
            "source_pdf_artifact_id": crf.json()["artifact_id"],
            "page_number": 1,
            "region": {"x": 0.1, "y": 0.2, "width": 0.4, "height": 0.1},
            "annotation_text": "Subject identifier",
            "links": [
                {
                    "target_type": "sdtm_variable",
                    "specification_revision": 1,
                    "domain": "DM",
                    "variable": "USUBJID",
                },
                {"target_type": "raw_variable", "raw_variable_id": raw_variable_id},
            ],
        },
    )
    assert annotation.status_code == 201, annotation.text
    linked = await client.post(
        f"/api/v1/acrf-annotations/{annotation.json()['id']}/links",
        headers={**headers, "If-Match": annotation.headers["etag"]},
        json={
            "operation": "amend",
            "links": [
                {
                    "target_type": "derivation_plan_node",
                    "plan_revision": 1,
                    "node_id": "derive",
                }
            ],
        },
    )
    assert linked.status_code == 200, linked.text
    assert len(linked.json()["links"]) == 3
    coverage = await client.get(f"/api/v1/sdtms/{sdtm.json()['id']}/acrf/coverage", headers=headers)
    assert coverage.status_code == 200, coverage.text
    assert coverage.json()["annotation_count"] == 1
    assert {tuple(item.values()) for item in coverage.json()["uncovered_sdtm_variables"]} == {
        ("DM", "STUDYID")
    }
    validation = await client.post(f"/api/v1/sdtms/{sdtm.json()['id']}/validate", headers=headers)
    assert validation.json()["status"] == "succeeded", validation.text

    metadata_artifact_types = (
        "sdtm_specification",
        "sdtm_derivation_report",
        "sdtm_lineage_report",
        "sdtm_validation_report",
    )
    with factory() as session:
        _, _, metadata_resolved = resolve_build_input(
            session,
            application.state.module_registry,
            study_id=UUID(study_id),
            owner_type="sdtm_workbench",
            owner_id=UUID(sdtm.json()["id"]),
            artifact_types=metadata_artifact_types,
            options={},
        )
    metadata_build = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/artifacts",
        headers={**headers, "Idempotency-Key": "build-sdtm-metadata-1"},
        json={
            "artifact_types": list(metadata_artifact_types),
            "expected_input_fingerprint": fingerprint(metadata_resolved),
        },
    )
    assert metadata_build.status_code == 202, metadata_build.text
    with factory.begin() as session:
        metadata_job = JobService().claim(session, worker_id="phase3-worker", lease_seconds=60)
        assert metadata_job is not None
    _process_job(
        factory,
        application.state.module_registry,
        application.state.object_storage,
        metadata_job,
    )
    completed_metadata = await client.get(
        f"/api/v1/artifact-builds/{metadata_build.json()['id']}", headers=headers
    )
    assert completed_metadata.json()["status"] == "succeeded", completed_metadata.text
    metadata_artifacts = [
        await client.get(f"/api/v1/artifacts/{artifact_id}", headers=headers)
        for artifact_id in completed_metadata.json()["result_artifact_ids"]
    ]
    assert {item.json()["artifact_type"] for item in metadata_artifacts} == set(
        metadata_artifact_types
    )
    specification_artifact = next(
        item for item in metadata_artifacts if item.json()["artifact_type"] == "sdtm_specification"
    )
    specification_download = await client.get(
        f"/api/v1/artifacts/{specification_artifact.json()['id']}/download", headers=headers
    )
    assert load_workbook(io.BytesIO(specification_download.content)).sheetnames == [
        "Metadata",
        "Domains",
        "Variables",
        "ValueLevel",
    ]

    with factory() as session:
        _, _, acrf_resolved = resolve_build_input(
            session,
            application.state.module_registry,
            study_id=UUID(study_id),
            owner_type="sdtm_workbench",
            owner_id=UUID(sdtm.json()["id"]),
            artifact_types=("sdtm_acrf",),
            options={},
        )
    acrf_build = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/acrf/artifacts",
        headers={**headers, "Idempotency-Key": "build-acrf-1"},
        json={"expected_input_fingerprint": fingerprint(acrf_resolved)},
    )
    assert acrf_build.status_code == 202, acrf_build.text
    with factory.begin() as session:
        acrf_job = JobService().claim(session, worker_id="phase3-worker", lease_seconds=60)
        assert acrf_job is not None
    _process_job(
        factory,
        application.state.module_registry,
        application.state.object_storage,
        acrf_job,
    )
    acrf_completed = await client.get(
        f"/api/v1/artifact-builds/{acrf_build.json()['id']}", headers=headers
    )
    assert acrf_completed.json()["status"] == "succeeded", acrf_completed.text
    acrf_artifact = await client.get(
        f"/api/v1/artifacts/{acrf_completed.json()['result_artifact_ids'][0]}", headers=headers
    )
    assert acrf_artifact.json()["artifact_type"] == "sdtm_acrf"
    assert "_source_object_key" not in acrf_artifact.text
    acrf_download = await client.get(
        f"/api/v1/artifacts/{acrf_artifact.json()['id']}/download", headers=headers
    )
    assert len(PdfReader(io.BytesIO(acrf_download.content)).pages) == 1

    with factory() as session:
        _, _, resolved = resolve_build_input(
            session,
            application.state.module_registry,
            study_id=UUID(study_id),
            owner_type="sdtm_workbench",
            owner_id=UUID(sdtm.json()["id"]),
            artifact_types=("sdtm_program_python",),
            options={},
        )
    compiled = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/compile",
        headers={**headers, "Idempotency-Key": "compile-python-1"},
        json={"backend": "python", "expected_input_fingerprint": fingerprint(resolved)},
    )
    assert compiled.status_code == 202, compiled.text
    with factory.begin() as session:
        job = JobService().claim(session, worker_id="phase3-worker", lease_seconds=60)
        assert job is not None
    _process_job(
        factory,
        application.state.module_registry,
        application.state.object_storage,
        job,
    )
    completed = await client.get(
        f"/api/v1/artifact-builds/{compiled.json()['id']}", headers=headers
    )
    assert completed.json()["status"] == "succeeded", completed.text
    artifact = await client.get(
        f"/api/v1/artifacts/{completed.json()['result_artifact_ids'][0]}", headers=headers
    )
    assert artifact.json()["source_plan_type"] == "sdtm_derivation_plan"
    assert artifact.json()["source_specification_revision"] == 1
    assert artifact.json()["compiler_fingerprint"]
    python_program = artifact.json()
    python_execution = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/execute",
        headers={**headers, "Idempotency-Key": "execute-python-1"},
        json={
            "language": "python",
            "program_artifact_id": python_program["id"],
            "input_artifact_ids": [source.json()["artifact_id"]],
            "validation_run_id": validation.json()["id"],
            "input_selection_revision": 1,
            "source_specification_revision": 1,
            "source_plan_revision": 1,
            "compiler_fingerprint": python_program["compiler_fingerprint"],
            "expected_input_fingerprint": python_program["input_fingerprint"],
            "limits": {
                "cpu": 1,
                "memory_mb": 512,
                "runtime_seconds": 60,
                "output_bytes": 1048576,
            },
        },
    )
    assert python_execution.status_code == 202, python_execution.text
    assert python_execution.json()["status"] == "queued"
    with factory.begin() as session:
        execution_job = JobService().claim(
            session,
            worker_id="phase3-execution-worker",
            lease_seconds=60,
            job_types=("execution",),
        )
        assert execution_job is not None
    process_execution_job(
        factory,
        application.state.module_registry,
        application.state.object_storage,
        HostPythonExecutor(),
        application.state.settings,
        execution_job.id,
    )
    executed = await client.get(
        f"/api/v1/execution-runs/{python_execution.json()['id']}", headers=headers
    )
    assert executed.json()["status"] == "succeeded", executed.text
    study_artifacts = await client.get(
        f"/api/v1/studies/{study_id}/artifacts", headers=headers, params={"limit": 200}
    )
    execution_log = next(
        item
        for item in study_artifacts.json()["items"]
        if item["artifact_type"] == "sdtm_execution_log"
        and item["execution_run_id"] == python_execution.json()["id"]
    )
    execution_log_download = await client.get(
        f"/api/v1/artifacts/{execution_log['id']}/download", headers=headers
    )
    assert execution_log_download.json()["status"] == "succeeded"
    assert execution_log_download.json()["language"] == "python"
    outputs = await client.get(
        f"/api/v1/execution-runs/{python_execution.json()['id']}/outputs", headers=headers
    )
    assert [item["artifact_type"] for item in outputs.json()["items"]] == [
        "sdtm_dataset_csv",
        "sdtm_dataset_xpt",
    ]
    xpt_artifact_id = outputs.json()["items"][1]["artifact_id"]
    xpt_download = await client.get(
        f"/api/v1/artifacts/{xpt_artifact_id}/download", headers=headers
    )
    xpt_frame, xpt_metadata = pyreadstat.read_xport(io.BytesIO(xpt_download.content))
    assert xpt_metadata.table_name == "DM"
    assert xpt_frame[["STUDYID", "USUBJID"]].to_dict(orient="records") == [
        {"STUDYID": "P3-001", "USUBJID": "01"},
        {"STUDYID": "P3-001", "USUBJID": "02"},
    ]

    current_sdtm = await client.get(f"/api/v1/sdtms/{sdtm.json()['id']}", headers=headers)
    define = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/define/initialize",
        headers={**headers, "If-Match": current_sdtm.headers["etag"]},
        json={
            "expected_define_revision": None,
            "study_oid": "STUDY.P3-001",
            "metadata_version_oid": "MDV.SDTM.1",
            "name": "Phase 3 SDTM Define",
            "standard_name": "SDTM-IG",
            "standard_version": "3.3",
            "output_artifact_ids": [xpt_artifact_id],
        },
    )
    assert define.status_code == 200, define.text
    assert define.json()["payload"]["datasets"][0]["leaf_id"] == "LF.DM"
    assert define.json()["payload"]["leaves"][0]["artifact_id"] == xpt_artifact_id
    manual_payload = define.json()["payload"]
    manual_payload["datasets"][0]["label"] = "Manually maintained DM label"
    manual_payload["datasets"][0]["items"][1]["origin"]["acrf_annotation_ids"] = [
        annotation.json()["id"]
    ]
    maintained = await client.put(
        f"/api/v1/sdtms/{sdtm.json()['id']}/define",
        headers={**headers, "If-Match": define.headers["etag"]},
        json=manual_payload,
    )
    assert maintained.status_code == 200, maintained.text
    reinitialized = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/define/initialize",
        headers={**headers, "If-Match": maintained.headers["etag"]},
        json={
            "expected_define_revision": 2,
            "study_oid": "IGNORED",
            "metadata_version_oid": "IGNORED",
            "name": "Ignored initializer value",
            "standard_name": "IGNORED",
            "standard_version": "0",
            "output_artifact_ids": [xpt_artifact_id],
        },
    )
    assert reinitialized.status_code == 200, reinitialized.text
    assert reinitialized.json()["revision"] == 3
    assert reinitialized.json()["payload"]["datasets"][0]["label"] == "Manually maintained DM label"
    define_validation = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/define/validate", headers=headers
    )
    assert define_validation.status_code == 200, define_validation.text
    assert define_validation.json()["error_count"] == 0
    xpt_after_define = await client.get(f"/api/v1/artifacts/{xpt_artifact_id}", headers=headers)
    assert xpt_after_define.json()["status"] == "available"
    reconciliation = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/define/reconcile",
        headers={**headers, "Idempotency-Key": "reconcile-sdtm-define-1"},
        json={
            "define_revision": reinitialized.json()["revision"],
            "expected_define_fingerprint": reinitialized.json()["fingerprint"],
        },
    )
    assert reconciliation.status_code == 202, reconciliation.text
    with factory.begin() as session:
        reconciliation_job = JobService().claim(
            session, worker_id="phase3-worker", lease_seconds=60
        )
        assert reconciliation_job is not None
    _process_job(
        factory,
        application.state.module_registry,
        application.state.object_storage,
        reconciliation_job,
    )
    reconciled = await client.get(
        f"/api/v1/reconciliation-runs/{reconciliation.json()['id']}", headers=headers
    )
    assert reconciled.json()["status"] == "succeeded", reconciled.text
    assert reconciled.json()["error_count"] == 0
    assert len(reconciled.json()["report_artifact_ids"]) == 1

    for version in ("2.0", "2.1"):
        rendered_define = await client.get(
            f"/api/v1/sdtms/{sdtm.json()['id']}/define/render",
            headers=headers,
            params={"version": version},
        )
        assert rendered_define.status_code == 200, rendered_define.text
        validate_define_xml(rendered_define.content, version)
    define_artifact_types = ("sdtm_define_xml_20", "sdtm_define_xml_21")
    with factory() as session:
        _, _, define_resolved = resolve_build_input(
            session,
            application.state.module_registry,
            study_id=UUID(study_id),
            owner_type="sdtm_workbench",
            owner_id=UUID(sdtm.json()["id"]),
            artifact_types=define_artifact_types,
            options={},
        )
    define_build = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/define/artifacts",
        headers={**headers, "Idempotency-Key": "build-define-xml-1"},
        json={
            "versions": ["2.0", "2.1"],
            "expected_input_fingerprint": fingerprint(define_resolved),
        },
    )
    assert define_build.status_code == 202, define_build.text
    with factory.begin() as session:
        define_job = JobService().claim(session, worker_id="phase3-worker", lease_seconds=60)
        assert define_job is not None
    _process_job(
        factory,
        application.state.module_registry,
        application.state.object_storage,
        define_job,
    )
    completed_define = await client.get(
        f"/api/v1/artifact-builds/{define_build.json()['id']}", headers=headers
    )
    assert completed_define.json()["status"] == "succeeded", completed_define.text
    define_artifacts = [
        await client.get(f"/api/v1/artifacts/{artifact_id}", headers=headers)
        for artifact_id in completed_define.json()["result_artifact_ids"]
    ]
    assert {item.json()["artifact_type"] for item in define_artifacts} == set(define_artifact_types)
    assert all(
        item.json()["validation_run_id"] == define_validation.json()["id"]
        and item.json()["reconciliation_run_id"] == reconciled.json()["id"]
        for item in define_artifacts
    )
    for item in define_artifacts:
        define_download = await client.get(
            f"/api/v1/artifacts/{item.json()['id']}/download", headers=headers
        )
        version = "2.0" if item.json()["artifact_type"].endswith("_20") else "2.1"
        validate_define_xml(define_download.content, version)

    current_sdtm = await client.get(f"/api/v1/sdtms/{sdtm.json()['id']}", headers=headers)
    csdrg = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/csdrg/initialize",
        headers={**headers, "If-Match": current_sdtm.headers["etag"]},
        json={
            "expected_csdrg_revision": None,
            "title": "Phase 3 Study Data Reviewer's Guide",
            "version": "1.0",
            "define_revision": reinitialized.json()["revision"],
            "reconciliation_run_id": reconciled.json()["id"],
        },
    )
    assert csdrg.status_code == 200, csdrg.text
    assert csdrg.json()["payload"]["dataset_facts"][0]["dataset"] == "DM"
    assert csdrg.json()["payload"]["dataset_facts"][0]["row_count"] == 2
    incomplete_csdrg_validation = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/csdrg/validate", headers=headers
    )
    assert incomplete_csdrg_validation.status_code == 200, incomplete_csdrg_validation.text
    assert incomplete_csdrg_validation.json()["error_count"] == 9

    csdrg_payload = csdrg.json()["payload"]
    csdrg_payload["narrative_sections"] = [
        {
            "section_id": section_id,
            "heading": heading,
            "paragraphs": [f"Manually authored {heading.lower()} content."],
        }
        for section_id, heading in (
            ("study_data_overview", "Study Data Overview"),
            ("standards_and_versions", "Standards and Versions"),
            ("trial_design", "Trial Design"),
            ("domain_notes", "Domain Notes"),
            ("conformance_issues", "Conformance Issues"),
            ("sponsor_defined_domains", "Sponsor-Defined Domains"),
            ("controlled_terminology", "Controlled Terminology"),
            ("data_dependencies", "Data Dependencies"),
            ("reviewer_guidance", "Reviewer Guidance"),
        )
    ]
    authored_csdrg = await client.put(
        f"/api/v1/sdtms/{sdtm.json()['id']}/csdrg",
        headers={**headers, "If-Match": csdrg.headers["etag"]},
        json=csdrg_payload,
    )
    assert authored_csdrg.status_code == 200, authored_csdrg.text
    refreshed_csdrg = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/csdrg/initialize",
        headers={**headers, "If-Match": authored_csdrg.headers["etag"]},
        json={
            "expected_csdrg_revision": 2,
            "title": "Ignored initializer title",
            "version": "ignored",
            "define_revision": reinitialized.json()["revision"],
            "reconciliation_run_id": reconciled.json()["id"],
        },
    )
    assert refreshed_csdrg.status_code == 200, refreshed_csdrg.text
    assert refreshed_csdrg.json()["payload"]["title"] == csdrg_payload["title"]
    assert (
        refreshed_csdrg.json()["payload"]["narrative_sections"]
        == csdrg_payload["narrative_sections"]
    )
    csdrg_validation = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/csdrg/validate", headers=headers
    )
    assert csdrg_validation.status_code == 200, csdrg_validation.text
    assert csdrg_validation.json()["error_count"] == 0

    with factory() as session:
        _, _, csdrg_resolved = resolve_build_input(
            session,
            application.state.module_registry,
            study_id=UUID(study_id),
            owner_type="sdtm_workbench",
            owner_id=UUID(sdtm.json()["id"]),
            artifact_types=("sdtm_csdrg_docx", "sdtm_csdrg_pdf"),
            options={},
        )
    csdrg_build = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/csdrg/artifacts",
        headers={**headers, "Idempotency-Key": "build-csdrg-1"},
        json={
            "artifact_types": ["sdtm_csdrg_docx", "sdtm_csdrg_pdf"],
            "expected_input_fingerprint": fingerprint(csdrg_resolved),
        },
    )
    assert csdrg_build.status_code == 202, csdrg_build.text
    with factory.begin() as session:
        csdrg_job = JobService().claim(session, worker_id="phase3-worker", lease_seconds=60)
        assert csdrg_job is not None
    _process_job(
        factory,
        application.state.module_registry,
        application.state.object_storage,
        csdrg_job,
    )
    completed_csdrg = await client.get(
        f"/api/v1/artifact-builds/{csdrg_build.json()['id']}", headers=headers
    )
    assert completed_csdrg.json()["status"] == "succeeded", completed_csdrg.text
    csdrg_artifacts = [
        await client.get(f"/api/v1/artifacts/{artifact_id}", headers=headers)
        for artifact_id in completed_csdrg.json()["result_artifact_ids"]
    ]
    assert all(
        item.json()["reconciliation_run_id"] == reconciled.json()["id"] for item in csdrg_artifacts
    )
    csdrg_docx = await client.get(
        f"/api/v1/artifacts/{csdrg_artifacts[0].json()['id']}/download", headers=headers
    )
    document = DocxDocument(io.BytesIO(csdrg_docx.content))
    assert any("Manually authored" in item.text for item in document.paragraphs)
    csdrg_pdf = await client.get(
        f"/api/v1/artifacts/{csdrg_artifacts[1].json()['id']}/download", headers=headers
    )
    assert len(PdfReader(io.BytesIO(csdrg_pdf.content)).pages) >= 2

    with factory() as session:
        _, _, sas_resolved = resolve_build_input(
            session,
            application.state.module_registry,
            study_id=UUID(study_id),
            owner_type="sdtm_workbench",
            owner_id=UUID(sdtm.json()["id"]),
            artifact_types=("sdtm_program_sas",),
            options={},
        )
    sas_build = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/compile",
        headers={**headers, "Idempotency-Key": "compile-sas-1"},
        json={"backend": "sas", "expected_input_fingerprint": fingerprint(sas_resolved)},
    )
    with factory.begin() as session:
        sas_job = JobService().claim(session, worker_id="phase3-worker", lease_seconds=60)
        assert sas_job is not None
    _process_job(
        factory,
        application.state.module_registry,
        application.state.object_storage,
        sas_job,
    )
    sas_completed = await client.get(
        f"/api/v1/artifact-builds/{sas_build.json()['id']}", headers=headers
    )
    sas_artifact = await client.get(
        f"/api/v1/artifacts/{sas_completed.json()['result_artifact_ids'][0]}", headers=headers
    )
    sas_program = sas_artifact.json()
    execution = await client.post(
        f"/api/v1/sdtms/{sdtm.json()['id']}/execute",
        headers={**headers, "Idempotency-Key": "execute-sas-1"},
        json={
            "language": "sas",
            "program_artifact_id": sas_program["id"],
            "input_artifact_ids": [source.json()["artifact_id"]],
            "validation_run_id": validation.json()["id"],
            "input_selection_revision": 1,
            "source_specification_revision": 1,
            "source_plan_revision": 1,
            "compiler_fingerprint": sas_program["compiler_fingerprint"],
            "expected_input_fingerprint": sas_program["input_fingerprint"],
            "limits": {
                "cpu": 1,
                "memory_mb": 512,
                "runtime_seconds": 60,
                "output_bytes": 1048576,
            },
        },
    )
    assert execution.status_code == 202, execution.text
    assert execution.json()["status"] == "failed"
    assert execution.json()["failure"]["code"] == "runner_not_configured"

    removed = await client.delete(
        f"/api/v1/acrf-annotations/{annotation.json()['id']}",
        headers={**headers, "If-Match": linked.headers["etag"]},
    )
    assert removed.status_code == 204, removed.text
    annotations = await client.get(
        f"/api/v1/sdtms/{sdtm.json()['id']}/acrf/annotations", headers=headers
    )
    assert annotations.json() == []

    variable = initialized.json()["variables"][0]
    current_variable = await client.get(f"/api/v1/raw-variables/{variable['id']}", headers=headers)
    changed = await client.patch(
        f"/api/v1/raw-variables/{variable['id']}",
        headers={**headers, "If-Match": current_variable.headers["etag"]},
        json={"label": "Subject Identifier"},
    )
    assert changed.status_code == 200, changed.text
    stale_selection = await client.get(
        f"/api/v1/sdtms/{sdtm.json()['id']}/input-selection", headers=headers
    )
    assert stale_selection.json()["state"] == "input_refresh_required"
    stale_validation = await client.get(
        f"/api/v1/validation-runs/{validation.json()['id']}", headers=headers
    )
    assert stale_validation.json()["status"] == "stale"
    stale_execution = await client.get(
        f"/api/v1/execution-runs/{python_execution.json()['id']}", headers=headers
    )
    assert stale_execution.json()["status"] == "stale"
    blocked = await client.post(f"/api/v1/sdtms/{sdtm.json()['id']}/validate", headers=headers)
    assert blocked.status_code == 409
    assert blocked.json()["code"] == "sdtm_plan_misaligned"
