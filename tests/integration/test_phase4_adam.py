"""PostgreSQL-backed exact SDTM-to-ADaM authoring workflow."""

import io
import json
import os
import subprocess
import sys
from collections.abc import AsyncIterator
from datetime import timedelta
from uuid import UUID

import pyreadstat
import pytest
from docx import Document as DocxDocument
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import text

from app.adam.define_xml import validate_define_xml
from app.artifacts.builds import resolve_build_input
from app.artifacts.models import Artifact
from app.catalog_seed import seed
from app.common.canonical import fingerprint
from app.common.time import utc_now
from app.database import create_database_engine, create_session_factory
from app.documents.service import stable_entity_id
from app.execution.adapters import ExecutionRequest, ExecutionResult
from app.execution.jobs import JobService
from app.execution.worker import process_execution_job
from app.main import create_app
from app.sdtm.contracts import SdtmSpecification
from app.sdtm.models import (
    SdtmDerivationPlanRevision,
    SdtmInputSelection,
    SdtmSpecificationRevision,
)
from app.sdtm.xpt import build_xpt_v5
from app.validation.models import ValidationRun
from tests.conftest import make_test_settings
from tests.integration.test_phase1_core import MemoryStorage, TestTokenVerifier
from workers.job_worker import _process_job

pytestmark = pytest.mark.skipif(
    os.getenv("AURINARA_INTEGRATION") != "1",
    reason="set AURINARA_INTEGRATION=1 with the Compose PostgreSQL service running",
)


class HostPythonExecutor:
    """Test-only adapter for executing deterministic generated Python."""

    def execute(self, request: ExecutionRequest) -> ExecutionResult:
        completed = subprocess.run(
            [sys.executable, "-I", str(request.workspace / "program.py")],
            check=False,
            capture_output=True,
            env={**os.environ, "AURINARA_WORKSPACE": str(request.workspace)},
            timeout=request.limits.runtime_seconds,
        )
        return ExecutionResult(
            status="succeeded" if completed.returncode == 0 else "failed",
            exit_code=completed.returncode,
            stdout=completed.stdout,
            stderr=completed.stderr,
            failure=None if completed.returncode == 0 else {"code": "program_failed"},
        )


@pytest.fixture
async def phase4_client() -> AsyncIterator[AsyncClient]:
    settings = make_test_settings()
    engine = create_database_engine(settings)
    factory = create_session_factory(engine)
    with engine.begin() as connection:
        connection.execute(text("TRUNCATE studies CASCADE"))
    storage = MemoryStorage()
    with factory() as session:
        seed(session, storage)
    app = create_app(
        settings=settings,
        session_factory=factory,
        object_storage=storage,
        token_verifier=TestTokenVerifier(),
    )
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://testserver"
    ) as client:
        yield client
    engine.dispose()


@pytest.mark.anyio
async def test_exact_sdtm_selection_and_adam_plan_alignment(
    phase4_client: AsyncClient,
) -> None:
    client = phase4_client
    headers = {"Authorization": "Bearer test-token"}
    study = await client.post(
        "/api/v1/studies",
        headers=headers,
        json={"code": "P4-001", "name": "Phase 4"},
    )
    study_id = study.json()["id"]
    sdtm = await client.post(
        f"/api/v1/studies/{study_id}/sdtms",
        headers=headers,
        json={"code": "SDTM", "name": "SDTM"},
    )
    assert sdtm.status_code == 201, sdtm.text
    application = client._transport.app  # type: ignore[attr-defined]
    factory = application.state.session_factory
    sdtm_id = sdtm.json()["id"]
    study_uuid = UUID(study_id)
    sdtm_uuid = UUID(sdtm_id)
    dependencies = await client.get(f"/api/v1/sdtms/{sdtm_id}/dependencies", headers=headers)
    dependency_fingerprint = dependencies.json()["fingerprint"]
    source_payload = {"raw_id": "00000000-0000-0000-0000-000000000001"}
    specification_payload = {
        "schema_version": "1.0.0",
        "domains": [
            {
                "code": "DM",
                "label": "Demographics",
                "domain_class": "Special Purpose",
                "structure": "One record per subject",
                "keys": ["USUBJID"],
                "sort_order": ["USUBJID"],
                "expected_output_name": "dm.xpt",
                "variables": [
                    {
                        "name": "USUBJID",
                        "label": "Unique Subject Identifier",
                        "data_type": "character",
                        "length": 40,
                        "role": "Identifier",
                        "core": "Req",
                        "origin": "Assigned",
                    },
                    {
                        "name": "ARM",
                        "label": "Description of Planned Arm",
                        "data_type": "character",
                        "length": 20,
                        "role": "Topic",
                        "core": "Exp",
                        "origin": "Assigned",
                    },
                ],
            }
        ],
    }
    plan_payload = {
        "schema_version": "1.0.0",
        "specification_revision": 1,
        "input_selection_revision": 1,
        "nodes": [{"node_type": "raw_read", "id": "read", "raw_dataset_code": "DEMOG"}],
        "edges": [],
    }
    sdtm_specification = SdtmSpecification.model_validate(specification_payload)
    xpt = build_xpt_v5(
        b"USUBJID,ARM\n01,Placebo\n02,Active\n",
        sdtm_specification.domains[0],
    )
    xpt_checksum = str(xpt.manifest["sha256"])
    application.state.object_storage.objects["test/dm.xpt"] = xpt.data
    with factory.begin() as session:
        source = SdtmInputSelection(
            aggregate_id=sdtm_uuid,
            study_id=study_uuid,
            revision=1,
            state="current",
            payload=source_payload,
            fingerprint=fingerprint(source_payload),
        )
        specification = SdtmSpecificationRevision(
            aggregate_id=sdtm_uuid,
            study_id=study_uuid,
            revision=1,
            schema_version="1.0.0",
            payload=specification_payload,
            fingerprint=fingerprint(specification_payload),
            created_by="user-1",
        )
        plan = SdtmDerivationPlanRevision(
            aggregate_id=sdtm_uuid,
            study_id=study_uuid,
            revision=1,
            schema_version="1.0.0",
            specification_revision=1,
            input_selection_revision=1,
            state="aligned",
            payload=plan_payload,
            fingerprint=fingerprint(plan_payload),
            created_by="user-1",
        )
        artifact = Artifact(
            study_id=study_uuid,
            module_key="sdtm",
            artifact_type="sdtm_dataset_xpt",
            owner_type="sdtm_workbench",
            owner_id=sdtm_uuid,
            filename="dm.xpt",
            media_type="application/x-sas-xport",
            object_key="test/dm.xpt",
            sha256=xpt_checksum,
            size=len(xpt.data),
            input_fingerprint="d" * 64,
            dependency_selection_revision=1,
            dependency_fingerprint=dependency_fingerprint,
            input_selection_revision=1,
            source_specification_revision=1,
            source_specification_fingerprint=specification.fingerprint,
            source_plan_type="sdtm_derivation_plan",
            source_plan_revision=1,
            generator_name="phase4-test",
            generator_version="1.0.0",
            status="available",
            manifest=xpt.manifest,
            created_by="user-1",
        )
        validation = ValidationRun(
            study_id=study_uuid,
            module_key="sdtm",
            owner_type="sdtm_workbench",
            owner_id=sdtm_uuid,
            status="succeeded",
            validator_name="sdtm-authoring-validator",
            validator_version="1.0.0",
            input_fingerprint="e" * 64,
            input_manifest={
                "sdtm_revision": 1,
                "dependency_fingerprint": dependency_fingerprint,
                "input_selection_revision": 1,
                "specification_revision": 1,
                "plan_revision": 1,
            },
            started_at=utc_now(),
            completed_at=utc_now(),
            error_count=0,
            warning_count=0,
            info_count=0,
            issues=[],
            report_artifact_ids=[],
            created_by="user-1",
        )
        session.add_all([source, specification, plan, artifact, validation])
        session.flush()
        artifact_id = str(artifact.id)
        validation_id = str(validation.id)

    adam = await client.post(
        f"/api/v1/studies/{study_id}/adams",
        headers=headers,
        json={"code": "ADAM", "name": "ADaM"},
    )
    assert adam.status_code == 201, adam.text
    adam_id = adam.json()["id"]
    adam_dependencies = await client.put(
        f"/api/v1/adams/{adam_id}/dependencies",
        headers={**headers, "If-Match": adam.headers["etag"]},
        json={
            "templates": [
                {
                    "purpose": "presentation",
                    "artifact_type": artifact_type,
                    "template_version_id": "30000000-0000-0000-0000-000000000005",
                }
                for artifact_type in ("adam_adrg_docx", "adam_adrg_pdf")
            ]
        },
    )
    assert adam_dependencies.status_code == 200, adam_dependencies.text
    selected = await client.put(
        f"/api/v1/adams/{adam_id}/input-selection",
        headers={**headers, "If-Match": adam_dependencies.headers["etag"]},
        json={
            "sdtm_id": sdtm_id,
            "sdtm_revision": 1,
            "validation_run_id": validation_id,
            "sdtm_dependency_fingerprint": dependency_fingerprint,
            "xpt_artifacts": [
                {
                    "dataset_code": "DM",
                    "artifact_id": artifact_id,
                    "checksum": xpt_checksum,
                }
            ],
        },
    )
    assert selected.status_code == 200, selected.text
    adam_specification = {
        "datasets": [
            {
                "id": "10000000-0000-0000-0000-000000000001",
                "code": "ADSL",
                "label": "Subject-Level Analysis Dataset",
                "dataset_class": "SUBJECT LEVEL ANALYSIS DATASET",
                "structure": "One record per subject",
                "keys": ["USUBJID"],
                "sort_order": ["USUBJID"],
                "expected_output_name": "adsl.xpt",
                "variables": [
                    {
                        "id": "20000000-0000-0000-0000-000000000001",
                        "name": "USUBJID",
                        "label": "Unique Subject Identifier",
                        "data_type": "character",
                        "length": 40,
                        "role": "Identifier",
                        "origin": {
                            "origin_type": "Predecessor",
                            "source_kind": "sdtm",
                            "source_dataset": "DM",
                            "source_variable": "USUBJID",
                        },
                        "output_node_id": "read_dm",
                    },
                    {
                        "id": "20000000-0000-0000-0000-000000000002",
                        "name": "TRT01P",
                        "label": "Planned Treatment for Period 01",
                        "data_type": "character",
                        "length": 20,
                        "role": "Treatment",
                        "origin": {"origin_type": "Derived", "description": "Copy ARM."},
                        "method_id": "MT.TRT01P",
                        "output_node_id": "derive_trt",
                    },
                ],
            }
        ],
        "methods": [
            {
                "id": "MT.TRT01P",
                "name": "Planned treatment",
                "description": "Copy the SDTM planned arm.",
                "plan_node_ids": ["derive_trt"],
            }
        ],
    }
    saved_specification = await client.put(
        f"/api/v1/adams/{adam_id}/specification",
        headers={**headers, "If-Match": selected.headers["etag"]},
        json=adam_specification,
    )
    assert saved_specification.status_code == 200, saved_specification.text
    adam_plan = {
        "specification_revision": 1,
        "input_selection_revision": 1,
        "nodes": [
            {"node_type": "sdtm_read", "id": "read_dm", "dataset": "DM"},
            {
                "node_type": "derive",
                "id": "derive_trt",
                "target_dataset": "ADSL",
                "method_id": "MT.TRT01P",
                "assignments": {"TRT01P": {"kind": "column", "dataset": "DM", "column": "ARM"}},
            },
            {
                "node_type": "select",
                "id": "select_adsl",
                "columns": ["USUBJID", "TRT01P"],
            },
            {"node_type": "dataset_write", "id": "write_adsl", "dataset": "ADSL"},
        ],
        "edges": [
            {"source_node": "read_dm", "target_node": "derive_trt"},
            {"source_node": "derive_trt", "target_node": "select_adsl"},
            {"source_node": "select_adsl", "target_node": "write_adsl"},
        ],
    }
    saved_plan = await client.put(
        f"/api/v1/adams/{adam_id}/derivation-plan",
        headers={**headers, "If-Match": saved_specification.headers["etag"]},
        json=adam_plan,
    )
    assert saved_plan.status_code == 200, saved_plan.text
    validated = await client.post(
        f"/api/v1/adams/{adam_id}/derivation-plan/validate", headers=headers
    )
    assert validated.status_code == 200, validated.text
    assert validated.json()["status"] == "valid"

    preview = await client.post(
        f"/api/v1/adams/{adam_id}/derivation-plan/preview",
        headers=headers,
        json={"datasets": ["ADSL"], "limit": 1},
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["outputs"]["ADSL"] == [{"USUBJID": "01", "TRT01P": "Placebo"}]
    lineage = await client.get(f"/api/v1/adams/{adam_id}/lineage", headers=headers)
    assert lineage.status_code == 200, lineage.text
    assert lineage.json()["outputs"]["ADSL"] == ["derive_trt", "read_dm", "select_adsl"]

    validation = await client.post(f"/api/v1/adams/{adam_id}/validate", headers=headers)
    assert validation.status_code == 200, validation.text
    assert validation.json()["status"] == "succeeded"
    assert validation.json()["error_count"] == 0

    metadata_artifact_types = (
        "adam_specification",
        "adam_specification_json",
        "adam_derivation_plan_json",
        "adam_derivation_report",
        "adam_traceability_report",
        "adam_validation_report",
        "adam_program_r",
        "adam_program_sas",
    )
    with factory() as session:
        _, _, metadata_resolved = resolve_build_input(
            session,
            application.state.module_registry,
            study_id=study_uuid,
            owner_type="adam_workbench",
            owner_id=UUID(adam_id),
            artifact_types=metadata_artifact_types,
            options={},
        )
    metadata_build = await client.post(
        f"/api/v1/adams/{adam_id}/artifacts",
        headers={**headers, "Idempotency-Key": "adam-metadata-artifacts-1"},
        json={
            "artifact_types": list(metadata_artifact_types),
            "expected_input_fingerprint": fingerprint(metadata_resolved),
        },
    )
    assert metadata_build.status_code == 202, metadata_build.text
    with factory.begin() as session:
        metadata_job = JobService().claim(
            session,
            worker_id="phase4-worker",
            lease_seconds=60,
            now=utc_now() + timedelta(seconds=1),
        )
        assert metadata_job is not None
    _process_job(
        factory,
        application.state.module_registry,
        application.state.object_storage,
        metadata_job,
    )
    completed_metadata = await client.get(
        f"/api/v1/artifact-builds/{metadata_build.json()['id']}", headers=headers
    )
    assert completed_metadata.json()["status"] == "succeeded", completed_metadata.text
    metadata_artifacts = [
        await client.get(f"/api/v1/artifacts/{item}", headers=headers)
        for item in completed_metadata.json()["result_artifact_ids"]
    ]
    assert {item.json()["artifact_type"] for item in metadata_artifacts} == set(
        metadata_artifact_types
    )

    with factory() as session:
        _, _, resolved = resolve_build_input(
            session,
            application.state.module_registry,
            study_id=study_uuid,
            owner_type="adam_workbench",
            owner_id=UUID(adam_id),
            artifact_types=("adam_program_python",),
            options={},
        )
    compiled = await client.post(
        f"/api/v1/adams/{adam_id}/compile",
        headers={**headers, "Idempotency-Key": "adam-compile-python-1"},
        json={"backend": "python", "expected_input_fingerprint": fingerprint(resolved)},
    )
    assert compiled.status_code == 202, compiled.text
    with factory.begin() as session:
        build_job = JobService().claim(
            session,
            worker_id="phase4-worker",
            lease_seconds=60,
            now=utc_now() + timedelta(seconds=1),
        )
        assert build_job is not None
    _process_job(
        factory,
        application.state.module_registry,
        application.state.object_storage,
        build_job,
    )
    completed_build = await client.get(
        f"/api/v1/artifact-builds/{compiled.json()['id']}", headers=headers
    )
    assert completed_build.json()["status"] == "succeeded", completed_build.text
    program = await client.get(
        f"/api/v1/artifacts/{completed_build.json()['result_artifact_ids'][0]}",
        headers=headers,
    )
    assert program.status_code == 200, program.text
    assert program.json()["artifact_type"] == "adam_program_python"
    assert program.json()["source_plan_type"] == "adam_derivation_plan"

    execution = await client.post(
        f"/api/v1/adams/{adam_id}/execute",
        headers={**headers, "Idempotency-Key": "adam-execute-python-1"},
        json={
            "language": "python",
            "program_artifact_id": program.json()["id"],
            "input_artifact_ids": [artifact_id],
            "validation_run_id": validation.json()["id"],
            "input_selection_revision": 1,
            "source_specification_revision": 1,
            "source_plan_revision": 1,
            "compiler_fingerprint": program.json()["compiler_fingerprint"],
            "expected_input_fingerprint": program.json()["input_fingerprint"],
            "limits": {
                "cpu": 1,
                "memory_mb": 512,
                "runtime_seconds": 60,
                "output_bytes": 1048576,
            },
        },
    )
    assert execution.status_code == 202, execution.text
    with factory.begin() as session:
        execution_job = JobService().claim(
            session,
            worker_id="phase4-execution-worker",
            lease_seconds=60,
            job_types=("execution",),
            now=utc_now() + timedelta(seconds=1),
        )
        assert execution_job is not None
    process_execution_job(
        factory,
        application.state.module_registry,
        application.state.object_storage,
        HostPythonExecutor(),
        application.state.settings,
        execution_job.id,
    )
    completed_execution = await client.get(
        f"/api/v1/execution-runs/{execution.json()['id']}", headers=headers
    )
    assert completed_execution.json()["status"] == "succeeded", completed_execution.text
    outputs = await client.get(
        f"/api/v1/execution-runs/{execution.json()['id']}/outputs", headers=headers
    )
    assert [item["artifact_type"] for item in outputs.json()["items"]] == [
        "adam_dataset_csv",
        "adam_dataset_xpt",
    ]
    xpt_output = await client.get(
        f"/api/v1/artifacts/{outputs.json()['items'][1]['artifact_id']}/download",
        headers=headers,
    )
    frame, metadata = pyreadstat.read_xport(io.BytesIO(xpt_output.content))
    assert metadata.table_name == "ADSL"
    assert frame.to_dict(orient="records") == [
        {"USUBJID": "01", "TRT01P": "Placebo"},
        {"USUBJID": "02", "TRT01P": "Active"},
    ]

    define = await client.post(
        f"/api/v1/adams/{adam_id}/define/initialize",
        headers={**headers, "If-Match": saved_plan.headers["etag"]},
        json={
            "expected_define_revision": None,
            "study_oid": "STUDY.P4-001",
            "metadata_version_oid": "MDV.ADAM.1",
            "name": "Phase 4 ADaM Define",
            "standard_name": "ADaM-IG",
            "standard_version": "1.3",
            "output_artifact_ids": [
                item["artifact_id"]
                for item in outputs.json()["items"]
                if item["artifact_type"] == "adam_dataset_xpt"
            ],
        },
    )
    assert define.status_code == 200, define.text
    assert define.json()["payload"]["datasets"][0]["code"] == "ADSL"
    assert define.json()["payload"]["methods"][0]["id"] == "MT.TRT01P"
    maintained_payload = define.json()["payload"]
    maintained_payload["description"] = "Manually maintained Define description"
    maintained = await client.put(
        f"/api/v1/adams/{adam_id}/define",
        headers={**headers, "If-Match": define.headers["etag"]},
        json=maintained_payload,
    )
    assert maintained.status_code == 200, maintained.text
    reinitialized = await client.post(
        f"/api/v1/adams/{adam_id}/define/initialize",
        headers={**headers, "If-Match": maintained.headers["etag"]},
        json={
            "expected_define_revision": 2,
            "study_oid": "IGNORED",
            "metadata_version_oid": "IGNORED",
            "name": "Ignored",
            "standard_name": "IGNORED",
            "standard_version": "0",
            "output_artifact_ids": maintained_payload["output_artifact_ids"],
        },
    )
    assert reinitialized.status_code == 200, reinitialized.text
    assert reinitialized.json()["payload"]["description"] == (
        "Manually maintained Define description"
    )
    valid_define_payload = reinitialized.json()["payload"]
    invalid_define_payload = json.loads(json.dumps(valid_define_payload))
    invalid_define_payload["datasets"][0]["label"] = "Metadata that differs from Specification"
    invalid_define = await client.put(
        f"/api/v1/adams/{adam_id}/define",
        headers={**headers, "If-Match": reinitialized.headers["etag"]},
        json=invalid_define_payload,
    )
    assert invalid_define.status_code == 200, invalid_define.text
    invalid_define_validation = await client.post(
        f"/api/v1/adams/{adam_id}/define/validate", headers=headers
    )
    assert invalid_define_validation.status_code == 200, invalid_define_validation.text
    assert invalid_define_validation.json()["error_count"] == 1
    assert invalid_define_validation.json()["issues"][0]["code"] == (
        "adam_define_dataset_metadata_mismatch"
    )
    blocked_render = await client.get(
        f"/api/v1/adams/{adam_id}/define/render",
        headers=headers,
        params={"version": "2.1"},
    )
    assert blocked_render.status_code == 409, blocked_render.text
    reinitialized = await client.put(
        f"/api/v1/adams/{adam_id}/define",
        headers={**headers, "If-Match": invalid_define.headers["etag"]},
        json=valid_define_payload,
    )
    assert reinitialized.status_code == 200, reinitialized.text
    define_validation = await client.post(
        f"/api/v1/adams/{adam_id}/define/validate", headers=headers
    )
    assert define_validation.status_code == 200, define_validation.text
    assert define_validation.json()["error_count"] == 0
    assert define_validation.json()["validator_version"] == "1.2.0"
    for version in ("2.0", "2.1"):
        rendered = await client.get(
            f"/api/v1/adams/{adam_id}/define/render",
            headers=headers,
            params={"version": version},
        )
        assert rendered.status_code == 200, rendered.text
        validate_define_xml(rendered.content, version)
    reconciliation = await client.post(
        f"/api/v1/adams/{adam_id}/define/reconcile",
        headers={**headers, "Idempotency-Key": "adam-define-reconcile-1"},
        json={
            "define_revision": reinitialized.json()["revision"],
            "expected_define_fingerprint": reinitialized.json()["fingerprint"],
        },
    )
    assert reconciliation.status_code == 202, reconciliation.text
    with factory.begin() as session:
        reconciliation_job = JobService().claim(
            session,
            worker_id="phase4-worker",
            lease_seconds=60,
            now=utc_now() + timedelta(seconds=1),
        )
        assert reconciliation_job is not None
    _process_job(
        factory,
        application.state.module_registry,
        application.state.object_storage,
        reconciliation_job,
    )
    completed_reconciliation = await client.get(
        f"/api/v1/reconciliation-runs/{reconciliation.json()['id']}", headers=headers
    )
    assert completed_reconciliation.json()["status"] == "succeeded"
    assert completed_reconciliation.json()["error_count"] == 0
    define_artifact_types = ("adam_define_xml_20", "adam_define_xml_21")
    with factory() as session:
        _, _, define_resolved = resolve_build_input(
            session,
            application.state.module_registry,
            study_id=study_uuid,
            owner_type="adam_workbench",
            owner_id=UUID(adam_id),
            artifact_types=define_artifact_types,
            options={},
        )
    define_build = await client.post(
        f"/api/v1/adams/{adam_id}/define/artifacts",
        headers={**headers, "Idempotency-Key": "adam-define-artifacts-1"},
        json={
            "versions": ["2.0", "2.1"],
            "expected_input_fingerprint": fingerprint(define_resolved),
        },
    )
    assert define_build.status_code == 202, define_build.text
    with factory.begin() as session:
        define_job = JobService().claim(
            session,
            worker_id="phase4-worker",
            lease_seconds=60,
            now=utc_now() + timedelta(seconds=1),
        )
        assert define_job is not None
    _process_job(
        factory,
        application.state.module_registry,
        application.state.object_storage,
        define_job,
    )
    completed_define = await client.get(
        f"/api/v1/artifact-builds/{define_build.json()['id']}", headers=headers
    )
    assert completed_define.json()["status"] == "succeeded", completed_define.text
    for artifact_id, version in zip(
        completed_define.json()["result_artifact_ids"], ("2.0", "2.1"), strict=True
    ):
        download = await client.get(f"/api/v1/artifacts/{artifact_id}/download", headers=headers)
        validate_define_xml(download.content, version)
        assert b'SourceSystemVersion="1.1.0"' in download.content

    adrg = await client.post(
        f"/api/v1/adams/{adam_id}/adrg/initialize",
        headers={**headers, "If-Match": reinitialized.headers["etag"]},
        json={
            "expected_adrg_revision": None,
            "title": "Phase 4 Analysis Data Reviewer's Guide",
            "version": "1.0",
            "define_revision": reinitialized.json()["revision"],
            "reconciliation_run_id": completed_reconciliation.json()["id"],
        },
    )
    assert adrg.status_code == 200, adrg.text
    assert adrg.json()["payload"]["dataset_facts"][0]["dataset"] == "ADSL"
    assert adrg.json()["payload"]["execution_facts"][0]["language"] == "python"
    adrg_payload = adrg.json()["payload"]
    for section in adrg_payload["narrative_sections"]:
        section["paragraphs"] = [
            f"Manually authored synthetic reviewer guidance for {section['heading']}."
        ]
    authored_adrg = await client.put(
        f"/api/v1/adams/{adam_id}/adrg",
        headers={**headers, "If-Match": adrg.headers["etag"]},
        json=adrg_payload,
    )
    assert authored_adrg.status_code == 200, authored_adrg.text
    refreshed_adrg = await client.post(
        f"/api/v1/adams/{adam_id}/adrg/initialize",
        headers={**headers, "If-Match": authored_adrg.headers["etag"]},
        json={
            "expected_adrg_revision": 2,
            "title": "Ignored",
            "version": "0",
            "define_revision": reinitialized.json()["revision"],
            "reconciliation_run_id": completed_reconciliation.json()["id"],
        },
    )
    assert refreshed_adrg.status_code == 200, refreshed_adrg.text
    assert all(
        section["paragraphs"] for section in refreshed_adrg.json()["payload"]["narrative_sections"]
    )
    adrg_validation = await client.post(f"/api/v1/adams/{adam_id}/adrg/validate", headers=headers)
    assert adrg_validation.status_code == 200, adrg_validation.text
    assert adrg_validation.json()["error_count"] == 0
    adrg_artifact_types = ("adam_adrg_docx", "adam_adrg_pdf")
    with factory() as session:
        _, _, adrg_resolved = resolve_build_input(
            session,
            application.state.module_registry,
            study_id=study_uuid,
            owner_type="adam_workbench",
            owner_id=UUID(adam_id),
            artifact_types=adrg_artifact_types,
            options={},
        )
    adrg_build = await client.post(
        f"/api/v1/adams/{adam_id}/adrg/artifacts",
        headers={**headers, "Idempotency-Key": "adam-adrg-artifacts-1"},
        json={
            "artifact_types": list(adrg_artifact_types),
            "expected_input_fingerprint": fingerprint(adrg_resolved),
        },
    )
    assert adrg_build.status_code == 202, adrg_build.text
    with factory.begin() as session:
        adrg_job = JobService().claim(
            session,
            worker_id="phase4-worker",
            lease_seconds=60,
            now=utc_now() + timedelta(seconds=1),
        )
        assert adrg_job is not None
    _process_job(
        factory,
        application.state.module_registry,
        application.state.object_storage,
        adrg_job,
    )
    completed_adrg = await client.get(
        f"/api/v1/artifact-builds/{adrg_build.json()['id']}", headers=headers
    )
    assert completed_adrg.json()["status"] == "succeeded", completed_adrg.text
    rendered_text: dict[str, str] = {}
    for artifact_id in completed_adrg.json()["result_artifact_ids"]:
        metadata = await client.get(f"/api/v1/artifacts/{artifact_id}", headers=headers)
        download = await client.get(f"/api/v1/artifacts/{artifact_id}/download", headers=headers)
        if metadata.json()["artifact_type"] == "adam_adrg_docx":
            document = DocxDocument(io.BytesIO(download.content))
            rendered_text["docx"] = "\n".join(item.text for item in document.paragraphs)
        else:
            reader = PdfReader(io.BytesIO(download.content))
            rendered_text["pdf"] = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "Phase 4 Analysis Data Reviewer's Guide" in rendered_text["docx"]
    assert "Manually authored synthetic reviewer guidance" in rendered_text["docx"]
    assert "Phase 4 Analysis Data Reviewer's Guide" in rendered_text["pdf"]

    exported = await client.get(
        f"/api/v1/adams/{adam_id}/specification/export?format=xlsx", headers=headers
    )
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(io.BytesIO(exported.content))
    workbook["Datasets"]["C2"] = "Imported subject-level dataset"
    changed = io.BytesIO()
    workbook.save(changed)
    workbook.close()
    import_headers = {
        **headers,
        "If-Match": refreshed_adrg.headers["etag"],
        "Idempotency-Key": "adam-specification-xlsx-1",
        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }
    imported = await client.post(
        f"/api/v1/adams/{adam_id}/specification/import",
        headers=import_headers,
        content=changed.getvalue(),
    )
    assert imported.status_code == 200, imported.text
    assert imported.json()["revision"] == 2
    assert imported.json()["payload"]["datasets"][0]["label"] == ("Imported subject-level dataset")
    replay = await client.post(
        f"/api/v1/adams/{adam_id}/specification/import",
        headers=import_headers,
        content=changed.getvalue(),
    )
    assert replay.status_code == 200, replay.text
    assert replay.json()["id"] == imported.json()["id"]
    blocked = await client.post(
        f"/api/v1/adams/{adam_id}/derivation-plan/validate", headers=headers
    )
    assert blocked.status_code == 409
    assert blocked.json()["code"] == "adam_plan_misaligned"

    revised_sdtm = await client.put(
        f"/api/v1/sdtms/{sdtm_id}/specification",
        headers={**headers, "If-Match": sdtm.headers["etag"]},
        json=specification_payload,
    )
    assert revised_sdtm.status_code == 200, revised_sdtm.text
    stale_selection = await client.get(f"/api/v1/adams/{adam_id}/input-selection", headers=headers)
    assert stale_selection.status_code == 200, stale_selection.text
    assert stale_selection.json()["state"] == "input_refresh_required"
    with factory() as session:
        stale_define_validation = session.get(
            ValidationRun, UUID(define_validation.json()["id"])
        )
        assert stale_define_validation is not None
        assert stale_define_validation.status == "stale"


@pytest.mark.anyio
async def test_specification_replaces_verified_normalized_trace_targets(
    phase4_client: AsyncClient,
) -> None:
    client = phase4_client
    headers = {"Authorization": "Bearer test-token"}
    study = await client.post(
        "/api/v1/studies",
        headers=headers,
        json={"code": "P4-TRACE", "name": "Phase 4 Traceability"},
    )
    study_id = study.json()["id"]
    sap = await client.post(
        f"/api/v1/studies/{study_id}/documents",
        headers=headers,
        json={
            "document_code": "SAP-TRACE",
            "name": "Traceability SAP",
            "document_type": "statistical_analysis_plan",
        },
    )
    assert sap.status_code == 201, sap.text
    sap_id = sap.json()["id"]
    content = await client.put(
        f"/api/v1/documents/{sap_id}/content",
        headers={**headers, "If-Match": sap.headers["etag"]},
        json={
            "document_type": "statistical_analysis_plan",
            "title": "Traceability SAP",
            "version": "1",
            "protocol_alignment": [
                {
                    "id": "AL1",
                    "name": "Alignment",
                    "description": "Manual alignment.",
                    "protocol_entity_ids": ["OBJ1"],
                }
            ],
            "objectives": [{"id": "O1", "name": "Objective", "description": "Text"}],
            "endpoints": [{"id": "E1", "name": "Endpoint", "description": "Text"}],
            "estimands": [
                {
                    "id": "EST1",
                    "name": "Estimand",
                    "description": "Text",
                    "population_id": "POP1",
                    "treatment_condition": "A versus B",
                    "variable": "Change",
                    "intercurrent_event_strategy": "Treatment policy",
                    "summary_measure": "Difference",
                    "protocol_entity_ids": ["END1"],
                }
            ],
            "populations": [{"id": "POP1", "name": "ITT", "description": "Text"}],
            "methods": [
                {"id": "M1", "name": "ANCOVA", "description": "Manually authored method."}
            ],
            "missing_data": [{"id": "MD1", "name": "Missing", "description": "Text"}],
            "output_requirements": [
                {"id": "OUT1", "name": "Table", "description": "Text"}
            ],
            "qc_items": [{"id": "QC1", "name": "QC", "description": "Text"}],
            "narrative_sections": [
                {"section_id": "methods", "heading": "Methods", "paragraphs": ["Text"]}
            ],
        },
    )
    assert content.status_code == 200, content.text
    adam = await client.post(
        f"/api/v1/studies/{study_id}/adams",
        headers=headers,
        json={"code": "ADAM-TRACE", "name": "ADaM traceability"},
    )
    assert adam.status_code == 201, adam.text
    adam_id = adam.json()["id"]
    dataset_id = "10000000-0000-0000-0000-000000000099"
    variable_id = "20000000-0000-0000-0000-000000000099"
    method_entity_id = str(stable_entity_id(UUID(sap_id), "method", "M1"))
    population_entity_id = str(stable_entity_id(UUID(sap_id), "population", "POP1"))
    specification = {
        "schema_version": "1.0.0",
        "datasets": [
            {
                "id": dataset_id,
                "code": "ADSL",
                "label": "Subject-Level Analysis Dataset",
                "dataset_class": "SUBJECT LEVEL ANALYSIS DATASET",
                "structure": "One record per subject",
                "keys": ["USUBJID"],
                "sort_order": ["USUBJID"],
                "expected_output_name": "adsl.xpt",
                "variables": [
                    {
                        "id": variable_id,
                        "name": "USUBJID",
                        "label": "Unique Subject Identifier",
                        "data_type": "character",
                        "length": 40,
                        "role": "Identifier",
                        "origin": {"origin_type": "Assigned"},
                        "output_node_id": "write_adsl",
                    }
                ],
            }
        ],
        "trace_targets": [
            {
                "relationship_type": "analyzes",
                "source": {
                    "module": "documents",
                    "entity_type": "sap_method",
                    "entity_id": method_entity_id,
                },
                "target": {
                    "module": "adam",
                    "entity_type": "adam_dataset",
                    "entity_id": dataset_id,
                },
            },
            {
                "relationship_type": "analyzes",
                "source": {
                    "module": "documents",
                    "entity_type": "sap_population",
                    "entity_id": population_entity_id,
                },
                "target": {
                    "module": "adam",
                    "entity_type": "adam_variable",
                    "entity_id": variable_id,
                },
            },
        ],
    }
    saved = await client.put(
        f"/api/v1/adams/{adam_id}/specification",
        headers={**headers, "If-Match": adam.headers["etag"]},
        json=specification,
    )
    assert saved.status_code == 200, saved.text
    links = await client.get(f"/api/v1/studies/{study_id}/trace-links", headers=headers)
    assert links.status_code == 200, links.text
    assert {
        (
            item["relationship_type"],
            item["source_entity_type"],
            item["source_entity_id"],
            item["target_entity_type"],
            item["target_entity_id"],
        )
        for item in links.json()
    } == {
        ("analyzes", "sap_method", method_entity_id, "adam_dataset", dataset_id),
        ("analyzes", "sap_population", population_entity_id, "adam_variable", variable_id),
    }

    invalid = await client.put(
        f"/api/v1/adams/{adam_id}/specification",
        headers={**headers, "If-Match": saved.headers["etag"]},
        json={
            **specification,
            "trace_targets": [
                {
                    **specification["trace_targets"][0],
                    "source": {
                        **specification["trace_targets"][0]["source"],
                        "entity_id": "60000000-0000-0000-0000-000000000099",
                    },
                }
            ],
        },
    )
    assert invalid.status_code == 422, invalid.text
    assert invalid.json()["code"] == "trace_endpoint_invalid"

    replaced = await client.put(
        f"/api/v1/adams/{adam_id}/specification",
        headers={**headers, "If-Match": saved.headers["etag"]},
        json={**specification, "trace_targets": []},
    )
    assert replaced.status_code == 200, replaced.text
    assert replaced.json()["revision"] == 2
    links = await client.get(f"/api/v1/studies/{study_id}/trace-links", headers=headers)
    assert links.json() == []
