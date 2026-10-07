"""SDTM Specification XLSX exchange tests."""

import io
from uuid import UUID

import pytest
from openpyxl import load_workbook

from app.sdtm.contracts import SdtmSpecification
from app.sdtm.specification_workbook import (
    parse_specification_xlsx,
    render_specification_xlsx,
)


def _specification() -> SdtmSpecification:
    terminology_release_id = UUID("f5bf5168-b8c1-44b3-962e-59c39af721ca")
    return SdtmSpecification.model_validate(
        {
            "domains": [
                {
                    "code": "DM",
                    "label": "Demographics",
                    "domain_class": "Special Purpose",
                    "structure": "One record per subject",
                    "keys": ["STUDYID", "USUBJID"],
                    "sort_order": ["STUDYID", "USUBJID"],
                    "expected_output_name": "dm.xpt",
                    "variables": [
                        {
                            "name": "STUDYID",
                            "label": "Study Identifier",
                            "data_type": "character",
                            "length": 20,
                            "role": "Identifier",
                            "core": "Req",
                            "origin": "Assigned",
                        },
                        {
                            "name": "USUBJID",
                            "label": "Unique Subject Identifier",
                            "data_type": "character",
                            "length": 40,
                            "role": "Identifier",
                            "core": "Req",
                            "origin": "Derived",
                            "derived": True,
                        },
                        {
                            "name": "SEX",
                            "label": "Sex",
                            "data_type": "character",
                            "length": 1,
                            "role": "Qualifier",
                            "core": "Req",
                            "origin": "Collected",
                            "terminology_binding": {
                                "terminology_release_id": terminology_release_id,
                                "codelist_code": "C66731",
                            },
                        },
                    ],
                    "value_level_metadata": [
                        {
                            "id": "DM.SEX.M",
                            "variable": "SEX",
                            "condition": {
                                "kind": "binary",
                                "operator": "eq",
                                "left": {"kind": "column", "dataset": "DM", "column": "SEX"},
                                "right": {"kind": "literal", "value": "M"},
                            },
                            "label": "Male subjects",
                            "data_type": "character",
                            "length": 1,
                            "terminology_binding": {
                                "terminology_release_id": terminology_release_id,
                                "codelist_code": "C66731",
                                "term_code": "C20197",
                            },
                        }
                    ],
                }
            ]
        }
    )


def test_specification_xlsx_round_trip_preserves_typed_content() -> None:
    specification = _specification()

    data = render_specification_xlsx(specification)
    observed = parse_specification_xlsx(data)

    assert observed.model_dump(mode="json") == specification.model_dump(mode="json")
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    assert workbook.sheetnames == ["Metadata", "Domains", "Variables", "ValueLevel"]
    workbook.close()


def test_specification_xlsx_rejects_cross_sheet_domain_reference() -> None:
    workbook = load_workbook(io.BytesIO(render_specification_xlsx(_specification())))
    workbook["Variables"]["A2"] = "UNKNOWN"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()

    with pytest.raises(ValueError, match="reference declared domains"):
        parse_specification_xlsx(output.getvalue())
